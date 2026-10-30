from pathlib import Path
from copy import deepcopy
from html import escape
import json
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
QA=Path(__file__).resolve().parent
ASSETS=QA/'structure-assets'; ASSETS.mkdir(exist_ok=True)
SOURCE=ROOT/'final/昆山杯-潮生智能-商业计划书.docx'
OUTPUT=ROOT/'final/昆山杯-潮生智能-商业计划书-人机共播修订版.docx'
FONT='C:/Windows/Fonts/msyh.ttc'
BOLD='C:/Windows/Fonts/msyhbd.ttc'

class Diagram:
    def __init__(self,name,w,h):
        self.name=name; self.w=w; self.h=h
        self.im=Image.new('RGB',(w,h),'white'); self.draw=ImageDraw.Draw(self.im)
        self.svg=[f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" viewBox="0 0 {w} {h}"><rect width="100%" height="100%" fill="white"/>']
    def line(self,points,arrow=False):
        self.draw.line(points,fill='#7A858E',width=3)
        self.svg.append('<polyline points="'+' '.join(f'{x},{y}' for x,y in points)+'" fill="none" stroke="#7A858E" stroke-width="3"/>')
        if arrow:
            x,y=points[-1]; px,py=points[-2]
            if abs(x-px)>abs(y-py): tri=[(x,y),(x-14 if x>px else x+14,y-8),(x-14 if x>px else x+14,y+8)]
            else: tri=[(x,y),(x-8,y-14 if y>py else y+14),(x+8,y-14 if y>py else y+14)]
            self.draw.polygon(tri,fill='#7A858E')
            self.svg.append('<polygon points="'+' '.join(f'{a},{b}' for a,b in tri)+'" fill="#7A858E"/>')
    def text(self,x,y,text,size=30,bold=False,center=False):
        font=ImageFont.truetype(BOLD if bold else FONT,size)
        for i,s in enumerate(text.split('\n')):
            xx=x-self.draw.textlength(s,font=font)/2 if center else x
            yy=y+i*int(size*1.5)
            self.draw.text((xx,yy),s,font=font,fill='#111111')
            self.svg.append(f'<text x="{x if center else xx}" y="{yy+size}" font-family="Microsoft YaHei" font-size="{size}" font-weight="{700 if bold else 400}" text-anchor="{"middle" if center else "start"}" fill="#111111">{escape(s)}</text>')
    def box(self,x,y,w,h,title,body='',size=29):
        self.draw.rectangle((x,y,x+w,y+h),outline='#B8C1C8',width=3)
        self.svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="white" stroke="#B8C1C8" stroke-width="3"/>')
        self.text(x+w/2,y+20,title,size,True,True)
        if body:self.text(x+w/2,y+size+42,body,size-4,False,True)
    def save(self):
        self.im.save(ASSETS/(self.name+'.png'))
        (ASSETS/(self.name+'.svg')).write_text(''.join(self.svg)+'</svg>',encoding='utf-8')
        return ASSETS/(self.name+'.png')

def figures():
    d=Diagram('01-human-ai-content',1700,390)
    for x,title,body in [(10,'真人主动造梗','设题  反差  追问'),(355,'AI即兴回应','性格  接梗  选择'),(700,'任务产生结果','游戏行动  真实反馈'),(1045,'观众参与接续','预测  共创  回访')]:
        d.box(x,30,300,155,title,body,28)
        if x<1045:d.line([(x+300,108),(x+343,108)],True)
    d.line([(1195,185),(1195,270),(160,270),(160,190)],True)
    d.text(690,305,'下一场回收共同经历  场后提炼完整切片',28,False,True)
    d.save()
    d=Diagram('02-audience-view',1700,820)
    d.box(15,15,1670,790,'观众直播画面方案',size=31)
    d.box(55,100,1060,430,'真实游戏画面或互动主题','能理解当前任务  能看见行动结果',31)
    d.box(1150,100,490,255,'姬岁歌角色形象','声音  表情  嘴型\n岁与歌的可辨识表达',29)
    d.box(1150,385,490,145,'真人搭档呈现','声音或头像  出镜形式待定',27)
    d.box(55,565,1585,90,'字幕与当晚目标',size=27)
    d.text(65,690,'平台原生区域：弹幕  关注  礼物  会员入口',29)
    d.text(65,740,'品牌段落与观众建议按节目需要展示',27)
    d.save()
    d=Diagram('03-operator-view',1700,640)
    d.box(20,20,500,220,'节目与角色','任务卡  人格选择\n语音打断  临场话题',30)
    d.box(600,20,500,220,'平台与运行','双平台来源  消息队列\n语音状态  推流检查',30)
    d.box(1180,20,500,220,'记忆与复盘','候选记忆  来源确认\n时间戳  高光筛选',30)
    d.box(20,320,500,220,'游戏与模型','模型配置  游戏状态\n建议审核  停止接管',30)
    d.box(600,320,500,220,'商业与结算','交付清单  账单核对\n会员权益  素材授权',30)
    d.box(1180,320,500,220,'观众画面输出','OBS场景  画面裁切\n字幕与音频混合',30)
    d.text(850,585,'示意按运营工作流组织  商业和双平台模块仍需补齐',28,False,True)
    d.save()
    d=Diagram('04-stack-to-practice',1700,650)
    boxes=[(20,20,'输入','真人语音  平台弹幕\n游戏事实与任务'),(600,20,'理解与调度','ASR  模型路由\n人格与上下文'),(1180,20,'角色表达','中文回应  接梗\nTTS  表情与嘴型'),(1180,340,'观众画面','Live2D  字幕\nOBS场景与音频'),(600,340,'游戏行动','动作检查  执行\n结果反馈与口播'),(20,340,'记忆与复盘','SQLite  人工审核\n下一场目标')]
    for x,y,title,body in boxes:d.box(x,y,500,230,title,body,30)
    d.line([(520,135),(595,135)],True);d.line([(1100,135),(1175,135)],True)
    d.line([(1430,250),(1430,335)],True)
    d.line([(850,250),(850,335)],True)
    d.line([(650,340),(650,285),(700,285),(700,255)],True)
    d.line([(600,455),(525,455)],True)
    d.line([(270,340),(270,290),(600,290),(600,240)],True)
    d.text(850,600,'模型接入形成能力基础  连续节目与双平台交付另行验收',28,False,True)
    d.save()
    d=Diagram('05-team-organization',1700,980)
    d.box(580,10,540,140,'杨若朝  首席执行官','项目统筹  资源与关键决策',30)
    d.line([(850,150),(850,205),(420,205),(420,250)],True)
    d.line([(850,205),(1280,205),(1280,250)],True)
    d.box(110,250,620,145,'徐嘉盛  技术总监','模型  游戏执行  系统验收',30)
    d.box(970,250,620,145,'姬祥  运营总监','节目排期  团队协作  运营交付',30)
    d.line([(420,395),(420,425),(70,425),(70,647)])
    d.line([(1280,395),(1280,415),(930,415),(930,882)])
    for y in [512,647]: d.line([(70,y),(105,y)],True)
    for y in [492,622,752,882]: d.line([(930,y),(965,y)],True)
    left=[('李岱松  系统工程师','部署  工作台  测试'),('肖鹏希  音频工程师','语音模型  声音效果')]
    right=[('刘瑞麒  AI人格与内容负责人','角色表达  接梗企划  社群'),('张涵  市场与商务负责人','受众研究  增长  品牌合作'),('吴越  视觉美术负责人','原画  Live2D制作协作'),('张乐溪  视频制作负责人','高光切片  系列包装')]
    for j,(a,b) in enumerate(left):d.box(110,455+j*135,620,115,a,b,27)
    for j,(a,b) in enumerate(right):d.box(970,435+j*130,620,115,a,b,27)
    d.text(420,790,'技术与运营共同参与播前检查\n主持和导演按场次轮值',28,False,True)
    d.save()
    d=Diagram('06-business-model',1700,670)
    d.box(445,10,810,130,'真人与AI共同直播','中文互动  游戏任务  连续角色关系',31)
    d.line([(850,140),(850,190)],True)
    d.box(445,195,810,110,'B站与抖音直播加切片传播',size=30)
    d.line([(850,305),(850,350),(440,350),(440,400)],True)
    d.line([(850,350),(1260,350),(1260,400)],True)
    d.box(60,400,760,180,'观众端付费','礼物与付费留言  会员订阅\n角色周边与纪念内容',29)
    d.box(880,400,760,180,'企业端付费','品牌主题直播  游戏体验合作\n角色与内容授权',29)
    d.text(850,615,'收入回流到节目  角色资产  社群与制作',28,False,True)
    d.save()
    d=Diagram('07-dual-platform',1700,520)
    d.box(560,5,580,120,'一套人机节目与导演流程',size=30)
    d.line([(850,125),(850,175),(390,175),(390,220)],True);d.line([(850,175),(1310,175),(1310,220)],True)
    d.box(30,220,720,175,'B站直播','横屏游戏与连续栏目\n弹幕  会员  主账号视频',30)
    d.box(950,220,720,175,'抖音直播','竖屏或适配画面  快速情境入口\n平台弹幕  礼物  短片分发',30)
    d.text(850,450,'分别接入消息  分别核算收入  共享角色与节目事实',28,False,True)
    d.save()

figures()
doc=Document(SOURCE)
body=doc._element.body
for child in list(body):
    if child.tag!=qn('w:sectPr'):body.remove(child)
sec=doc.sections[0]
sec.top_margin=Cm(2.1);sec.bottom_margin=Cm(1.9)
sec.left_margin=Cm(2.15);sec.right_margin=Cm(2.15)
for name,size in [('Normal',10.5),('Title',28),('Subtitle',13),('Heading 1',17),('Heading 2',12)]:
    st=doc.styles[name];st.font.name='Microsoft YaHei';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    st.paragraph_format.line_spacing=1.18
    st.paragraph_format.space_after=Pt(7)
    if name.startswith('Heading'):
        st.paragraph_format.keep_with_next=True
        st.paragraph_format.space_before=Pt(11)
    for e in list(st.element.get_or_add_pPr()):
        if e.tag in [qn('w:pBdr'),qn('w:shd')]:st.element.get_or_add_pPr().remove(e)
for para in sec.header.paragraphs:
    for run in para.runs:run.text=run.text.replace('姬岁歌AI虚拟主播商业计划书','姬岁歌人机共播商业计划书')
for para in sec.footer.paragraphs:
    for run in para.runs:run.text=run.text.replace('2026年10月21日','2026年10月24日')

pages=[]
def start(title):
    if pages:doc.add_page_break()
    pages.append(title)
    return doc.add_paragraph(title,'Heading 1')
def p(text,bold=False):
    el=doc.add_paragraph(text)
    if bold:
        for r in el.runs:r.bold=True
    return el
def h(text):return doc.add_paragraph(text,'Heading 2')
def fig(name,caption,width=16.6):
    el=doc.add_paragraph(); el.paragraph_format.space_before=Pt(4);el.paragraph_format.space_after=Pt(3)
    r=el.add_run();r.add_picture(str(ASSETS/(name+'.png')),width=Cm(width))
    inline=r._r.xpath('.//wp:docPr')
    if inline:inline[0].set('descr',caption)
    cap=doc.add_paragraph(caption);cap.paragraph_format.space_after=Pt(8)
    for r in cap.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('555555')
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for j,w in enumerate(widths):t.columns[j].width=Cm(w)
    for j,s in enumerate(headers):t.rows[0].cells[j].text=s
    for row in rows:
        c=t.add_row().cells
        for j,s in enumerate(row):c[j].text=str(s)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if not ri:trpr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
            pr.append(b)
            if not ri:
                x=OxmlElement('w:shd');x.set(qn('w:fill'),'ECEFF1');pr.append(x)
            m=OxmlElement('w:tcMar')
            for side,v in [('top',90),('bottom',90),('left',90),('right',90)]:
                x=OxmlElement('w:'+side);x.set(qn('w:w'),str(v));x.set(qn('w:type'),'dxa');m.append(x)
            pr.append(m)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.space_before=Pt(0);pp.paragraph_format.line_spacing=1.1
                if ci:pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in pp.runs:r.font.size=Pt(9);r.bold=not ri;r.font.color.rgb=RGBColor(0,0,0)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
def link(text,url):
    pp=doc.add_paragraph();pp.paragraph_format.space_after=Pt(7)
    hyp=OxmlElement('w:hyperlink');hyp.set(qn('r:id'),pp.part.relate_to(url,RT.HYPERLINK,is_external=True))
    rr=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'21506B');pr.append(color);rr.append(pr)
    tt=OxmlElement('w:t');tt.text=text;rr.append(tt);hyp.append(rr);pp._p.append(hyp)
    return pp

# 1 Cover and executive summary
pages.append('封面与摘要')
p('2026 昆山杯清华大学创业大赛')
doc.add_paragraph('潮生智能商业计划书','Title')
doc.add_paragraph('姬岁歌真人与AI共同直播项目','Subtitle')
p('团队名称  潮生智能团队\n项目负责人  杨若朝\n参赛组别  种子组\n参赛赛道  第四赛道 AI for X 行业应用与科学创新')
h('项目摘要')
p('潮生智能以原创角色姬岁歌为核心，开发真人与AI共同直播的互动娱乐节目。真人主动设题、造梗并推动现场情节，AI以岁与歌两种人格表达参与中文对话、弹幕互动和真实游戏任务。观众持续观看的是双方怎样合作、争论、认错和兑现约定。')
p('项目拟在Bilibili与抖音双平台直播，以北京时间20:00至22:00为初期档期，每周五场。完整直播与高光切片共同获客，角色关系和观众共创推动回访，收入来自礼物与付费留言、会员、周边及品牌和游戏合作。双平台开播及真实事件接入作为上线交付，尚不宣称完成商业化运行。')
p('现阶段已有对话、人格与记忆、语音、Live2D舞台和导演工作台等工程基础，部分游戏任务有可核对记录。近期优先交付可观看的人机共播样例，补齐中文接梗评测、原创视觉声线与平台联调，再进入固定公开运营。')
p('计划争取10万元启动支持。基准测算为首个完整运营年度收入61,650元，日常运营与周边履约支出62,500元，另计25,000元启动投入。以上为目标预算；资金、真人工时和实际回款按月复盘。')

# 2 contents
start('目录')
for title,num in [('一 项目背景',3),('二 市场规模与竞品分析',4),('三 产品介绍',6),('四 团队介绍',10),('五 商业模式',11),('六 市场战略规划',17),('公开来源与工程证据',20),('后续演示提炼提纲',22)]:
    pp=p(f'{title}\t{num}');pp.paragraph_format.space_after=Pt(11)
h('商业模式阅读索引')
for title,num in [('收入结构与付费对象',11),('游戏选择与直播内容商品化',12),('打赏与付费留言的运营细节',13),('会员与周边的持续经营',14),('品牌直播与游戏合作的交付',15),('预算与首年收支测算',16)]:
    p(f'{title}\t{num}')
h('产品介绍阅读索引')
p('观众界面与人机节目体验  第6页\n后台调试与导演工作台  第7页\n技术栈怎样支撑实际节目  第8页\n人格化与中文接梗的研发验证  第9页')

# 3 background
start('一 项目背景')
p('我们面向喜欢游戏、角色成长和现场互动的中文观众，提供可持续追更的人机搭档节目。项目要解决的问题是：AI能回答问题之后，怎样继续制造值得看、值得分享，并能按期交付的内容。')
h('真人与AI共同直播的优势')
p('真人擅长根据现场气氛设题、追问和转场，能解释情境并承担真实反应；AI能够以稳定角色设定回应，并在真实任务中给出不同于搭档预期的选择。双方互相触发，使一个回答发展为有铺垫、有反转、有结果的节目。技术对节目效果的贡献要通过完整片段和观众回访检验。')
fig('01-human-ai-content','节目机制示意  真人主动创造情境  AI保留即兴回应  由结果产生可接续内容')
h('从人与人互动中保留可共情的部分')
p('人与人联动的吸引力常来自默契、分歧和互相理解。本项目保留这些可观察关系：真人会惊讶、着急、认错，也会照顾搭档的表达空间；AI通过好奇、坚持、承认判断偏差和记起约定呈现人格化。AI的情绪属于角色表达，能力与进度仍按事实说明。')
h('为什么需要持续投入本地化')
p('中文开口只是基础。一个梗可能同时依赖谐音、弹幕上下文、游戏局势、搭档关系和说话时机。项目需要长期筛选角色适用的表达，记录观众真正产生反应的片段，避免把网上流行语机械塞进回复。真人主动造梗为AI提供可理解的现场条件，降低空泛聊天的概率。')

# 4 market size
start('二 市场规模与竞品分析')
h('用少量数据说明机会的存在')
p('YouTube官方披露，2024年其300个虚拟创作者样本在视频、直播与Shorts中合计获得超过150亿次观看。这说明虚拟角色有跨内容形式的观看需求；样本包括多种虚拟创作者，不能直接等同于AI主播市场规模或本项目收入空间。[S1]')
p('Bilibili2026年第二季度业绩披露，月活跃用户约3.71亿。平台用户规模提供分发背景，实际可服务人群仍需从中文游戏、VUP及AI互动兴趣中筛选，不能把平台全部用户计算成姬岁歌的潜在付费观众。[S2]')
h('可服务人群与进入路径')
p('首批服务两类观众：一类愿意持续追看角色与搭档关系，另一类先被游戏反转和AI意外反应吸引。B站承接较完整的游戏过程、连载与社群，抖音直播和短片用更快的情境入口触达陌生观众；两端都要提供持续直播和原生互动。')
p('初期以校内外成年受试者和自愿进入直播的陌生观众验证需求。首轮30人按既往观看习惯分层，观察自由观看、角色与事件回忆以及七天回访。校内熟人前测单独记录，不并入正式陌生人验证，也不把口头好评当成实际留存。')
h('市场规模如何落实为经营目标')
p('短期先证明节目有人愿意看、会记住并再次回来，再扩大曝光。原增长测算为全年约200条重点切片、15万主阵地年末关注目标，属于进取目标。公开运营后使用两平台真实播放、关注和回访重估，不根据行业观看总量直接外推粉丝或付费率。')
p('WWW 2025的VUP观众研究为粉丝社区与付费行为提供外部参照，研究对象并非本项目的AI直播观众。它支持提前观察核心粉丝与持续参与，不能替代本项目受众测试。[S3]')

# 5 competitors
start('市场规模中的竞品分析')
h('Neuro的优势是角色与关系已经形成内容')
p('Neuro-sama及Evil Neuro已有鲜明人物关系，官方站覆盖直播、视频、音乐与商品入口。值得学习的是角色与真人创作者的长期互动，以及从直播延伸到作品和商品的经营方式。姬岁歌当前的节目积累、粉丝基础和跨媒介经营成熟度都需要追赶。[S4]')
h('中文本土化与直播时间是进入切口')
p('本项目选择原生中文语境、本地游戏话题与北京时间晚间互动。把一个既有国际角色迁入中文场景，需要同时维持声线、性格、说话习惯和粉丝熟悉的关系；直接翻译未必能保留原有笑点。我们将中文本土化视为持续的内容投入，以观众能否理解、愿意互动和回访评价效果；对手的具体语言能力以其实际节目为准。')
p('2026年10月24日查询的第三方排期样本中，10月24日18:00 UTC与10月25日19:00 UTC分别对应北京时间次日02:00与03:00。它们说明部分时段与国内晚间观看不重合，不能推为所有直播场次的固定安排；排期以官方Twitch临时公告为准。姬岁歌拟争取20:00至22:00的现场互动位置。[S5]')
h('Lumi与Nox说明国内技术门槛已经提高')
p('Lumi_Nox公开项目包含双AI实时语音、发言仲裁、观众与角色记忆以及游戏桥接。记忆和语音接口本身不足以构成我们的独有优势。姬岁歌应把真人主动造梗、双人格在同一角色中的表现，以及中文接梗的实际效果落实为可观看样片。[S6]')
h('和普通真人主播比较什么')
p('真人主播具有成熟表达和现场判断，AI独立直播需要解决节奏与稳定性，人机共播则增加策划、配合和技术维护成本。我们的优势假设是互动反差与连续角色关系能提高内容辨识和回访。内部将比较相近任务和时长的样片，以实际观看与制作投入判断收益，现阶段不宣称已全面胜过真人或AI主播。')

# 6 product audience
start('三 产品介绍 观众界面与节目体验')
p('观众首先接触的是完整节目。直播画面以真实游戏或互动主题为主体，配合姬岁歌形象、真人搭档声音或头像、字幕与当晚目标。弹幕、关注和礼物使用各平台自己的入口。真人是否露脸由节目安排决定。')
fig('02-audience-view','观众界面方案示意  用于表达信息分工  不是已经上线的完整产品截图')
h('一段完整互动怎样发生')
p('真人提出“先当队长还是先干活”的木剑挑战，AI回应后在实际游戏中采集和制作；观众预测结果，真人根据进度追问。成功后回收开场的玩笑，失败则一起核对缺少的材料。观众既能看见反差，也能理解任务的真实结果和下次约定。')
h('已有前台与待交付部分')
p('工程已有Live2D舞台、字幕、弹幕显示和语音控制；首页仍包含本地测试入口。overlay=1样式可隐藏控制区用于OBS叠加。完整的双平台画面适配、原创角色成品、真人搭档布局及平台原生交互联调仍需验收，不能把本地调试页面直接作为商业直播成品。')

# 7 product backstage
start('产品介绍 后台调试与导演工作台')
p('后台服务团队，观众端服务观看。导演工作台应让主持、人格运营、技术和值班运营围绕同一场节目协作，查看任务、消息、语音、游戏事实和高光记录；调试参数、错误日志及模型配置不进入观众直播画面。')
fig('03-operator-view','后台工作流示意  现有工程可作为基础  双平台与商业账单模块需补齐')
h('现有页面分别承担什么职责')
p('首页与OBS叠加层承担角色舞台与字幕；games.html组织游戏画面、模型方案、聊天及任务停止；studio.html提供候选记忆审核、大脑诊断、连续经历和直播运行；debug.html用于开发调试。这些页面在信息和权限上仍需进一步梳理，隐藏控制区不等于已经完成后台权限隔离。')
h('商业开播需要补齐的工作')
p('双平台消息标记来源，礼物事件核对去重、撤回和退款状态，观众建议经确认后再影响任务；播前检查观众画面只包含批准内容，播后按时间戳导出高光与交付记录。商业收入以平台结算账单为准，模型口播感谢不承担记账职责。')

# 8 product tech
start('产品介绍 技术栈迁移到实际节目')
fig('04-stack-to-practice','技术到节目流程示意  每个模块都对应观众可感知的行为与结果')
h('聊天与人物表达')
p('Node.js服务与网页舞台连接输入、调度和展示。SenseVoice语音识别支持把真人语音送入上下文，模型路由结合人格与现场任务生成回应；GPT-SoVITS或系统TTS负责朗读，Live2D与字幕呈现角色。实际使用的声线、模型及延迟必须按试播版本说明，模块存在不等于所有组合都已稳定。')
h('游戏与记忆形成连续内容')
p('Minecraft以DeepSeek等规划方案与Mineflayer执行衔接，动作经过程序检查并读取结果；STS2有工具动作与口播回执快照。SQLite保存经人工确认的共同事件，供下次接续。场上讲话依据真实游戏状态，准备、尝试和完成采用不同表达，避免让模型把计划说成成果。')
h('当前可演示与下一步验收')
p('可核对材料包括中文对话与角色状态模块、记忆审核入口、语音调度代码和部分实机任务。待验收内容包括连续中文接梗效果、原创视觉与授权声线整合、双平台真实事件与稳定播出。完整游戏通关、任意游戏适配、商业转化及长期留存不列为已实现成果。')

# 9 product localization
start('产品介绍 人格化与中文接梗验证')
h('AI和真人的人性化一面')
p('岁偏向好奇、观察和理解，歌偏向坚持、选择与守护。她们需要在场景中表现立场：好奇时追问理由，失误时承认偏差，承诺过的事在下一场有进展。真人也呈现自己的困惑、胜负欲和关心，允许与AI意见不同，再通过实际行动推进关系。')
h('为什么接梗是专门的本地化工作')
p('梗能否成立取决于共享语境、包袱方向与回应时机。中文谐音和游戏术语可能依赖语音识别，反话和观众调侃需要关系判断，老梗需要可靠来源和使用节制；同一句话由岁或歌说出还应有不同态度。大模型能生成中文只说明具备语言基础，不能据此证明会自然接梗。')
h('先用真人造梗提供可测试场景')
p('第一步准备谐音、认真误解、反向挑战、角色分歧和旧事回收等场景，同时放入普通对话作为对照。第二步给出当晚任务、观众来源和角色立场，由AI即兴回应。第三步让真人根据回答继续互动，并记录是否出现解释笑话、答非所问、过度顺从、重复老梗或错过接话时机。')
h('从样本到真实试播的验证')
p('先安排不少于30个中文场景，每种人格分别生成并由内容人员复核；保留未见过的场景检验迁移。再制作两段时长、画面与声音相近的样片，首轮30名成年受试者交叉观看。一次只比较一项变化，例如有无主动造梗或有无共同记忆，避免同时变动多项后无法解释结果。')
p('记录语境理解、人格一致、自然接续、首次可听回应等待与重复情况；受试者自由回忆具体片段，观察可选观看和七天回访。自动评分用于定位问题，真人反馈与实际行为决定是否保留。节目形成的有效梗经来源审核、复测后再复用，尚未形成的测评数据不填写成功率。')

# 10 team
start('四 团队介绍')
p('九人团队按项目职责设置执行、技术与运营岗位。首席执行官统筹资源与方向，技术总监负责系统实现和验收，运营总监负责节目排期与跨工作线交付；主持和导演按场次轮值。以下为本计划中的团队岗位安排，公司设立后的正式任职另行确认。')
fig('05-team-organization','团队组织与岗位分工  内容 音频  视觉和研发按同一场节目协作',16.6)
h('每场节目怎样协作')
p('播前由运营总监组织任务、主持、抛梗与交付审阅，技术总监确认依赖服务和接管流程；人格内容负责人检查角色回应，视觉与音频负责人核对舞台。播后视频负责人筛选高光，市场商务负责人汇总观看与合作数据，CEO决定下一轮投入。成员投入工时单独记录，避免把协作成本隐藏在现金预算之外。')

# 11 business overview
start('五 商业模式 收入结构与付费对象')
p('团队经营原创人机搭档IP及其内容。观众为喜欢的现场互动和持续参与付费，企业为接触目标观众、完成主题内容与使用角色资产付费。可持续收入依赖节目回访和交付质量，技术能力先服务内容生产。')
fig('06-business-model','商业模式结构  观众付费与企业付费分别核算  不将礼物面额当作主播实收')
h('观众为什么愿意支持')
p('打赏回应当场的有趣表现，付费留言提供平台允许的醒目表达，会员支持持续追看与明确权益，周边保存角色和共同事件。真人主动造梗、AI即兴接梗产生值得传播的片段；跨场任务与人物关系提供回来的理由，权益与商品把这种关系转成可交付的产品。')
h('企业为什么选择这档节目')
p('品牌可以把主题写进人机共同任务，让真人解释、AI选择、观众参与，形成直播段落与二次传播素材。竞争点应是可验证的受众、自然的节目表达和可靠履约；数据未形成时先做小范围样例，不承诺观看人数、带货量或转化结果。')

# 12 games business
start('商业模式 游戏选择与内容商品化')
h('选择游戏先看能否形成可卖的内容')
p('优先选择有稳定观看需求、能制造直播效果、AI当前能够执行、观众容易参与的游戏类型。流量潜力决定新观众机会，明确任务和失败反馈决定切片效果，接入与维护成本决定能否持续播出。热度只是筛选条件，不能弥补长期卡顿和缺少互动。')
h('开放世界用于成长连载')
p('Minecraft适合共同目标、建造约定和观众建议，场景可延续到下一场，名场面与成果也可用于纪念内容。现有工程有采木与木剑制作等记录；复杂建家、生存稳定性仍需验证。初期选择材料与时限明确的短任务，把等待和恢复控制在可主持范围内，再扩成长线栏目。')
h('策略与回合制用于分歧和预测')
p('以STS2等已接入原型为起点，让真人与AI先判断出牌或路线，邀请观众预测，再由结果回收分歧。回合边界便于讨论和剪辑，适合主题体验合作。18次成功动作与口播结束回执仅证明局部链路，完整场次和游戏推广能力需要另行验收。')
h('轻量挑战用于短片与节奏补充')
p('Flappy Bird已有约定条件下的独立策略验收，适合短挑战和榜单话题，但不能据此声称已与人格语音完整整合。轻量游戏可填补节目节奏，不能只展示分数而缺少人机交流。反应与多人派对类列入候选，待适配、现场效果与游戏使用要求评估后安排。')
h('流量与效果怎样验证')
p('公开运营前对两平台相同观察窗口采集候选游戏的活跃直播、可获取观看、短片中位表现与同量级创作者更新情况，标明日期。内部同长度试播比较留存、有效互动、每场可用切片及制作工时。优先保留既有观看需求又能讲出完整故事的栏目，再安排投流和商务样片。')

# 13 gifts details
start('商业模式 打赏与付费留言细节')
h('付费入口与收入范围')
p('B站侧规划使用平台提供的礼物、付费留言及大航海等可开通能力；抖音侧使用其直播礼物及账号实际可用权益。会员类收入与普通礼物分开记录。具体价格、分成、开通条件与提现周期以账号当期规则和合同为准，不设置一个跨平台通用的“充值即到账”比例。[S7][S8]')
h('收到礼物后怎样回应')
p('礼物事件先核对来源与重复状态，再进入感谢队列。AI按角色风格作简短感谢，真人结合节目节奏接话；连续小额事件可合并，在任务间隙集中回应。付费留言先检查能否展示和回答，再在合适节点回应。关键游戏动作与已有发言受保护，礼物不能无限打断节目。')
p('感谢可以关联真实片段，例如“刚才你看见我差点翻车了”。回应不承诺私下关系，不把打赏绑定角色的爱意、痛苦或输赢，也不以消极游戏操作诱导送礼。公开任务与基本互动应让未付费观众也能参与，商业权益通过清晰规则交付。[S9]')
h('技术事件和平台账单分别处理')
p('现有B站适配层可归一化弹幕、礼物、付费留言与大航海事件，但不能据模拟事件证明真实直播间已接通。当前归一化结果以类型和文本为主，事件唯一标识、金额、订单生命周期、去重与退款对账需补齐；抖音真实事件接入也列为待交付。')
p('上线验收至少覆盖重复事件只回应一次、合并礼物不漏记、跨平台来源可辨、异常重连可追溯及账单可核对。感谢队列的显示值只用于节目，收入由平台结算单确认；退回、冻结或未结算金额不能作为可用现金。')
h('打赏测算怎样落到经营')
p('按场记录有效观看、实际支持人数、净结算收入和感谢造成的中断。内部测算可以使用“有效观众×实际付费比例×人均净收入”，各参数都来自同一平台、同一观察期。第一阶段检验是否出现自愿重复支持，不以礼物流水替代收入，也不把舰长或订阅再计入普通打赏。')

# 14 membership merch
start('商业模式 会员与周边的持续经营')
h('会员提供可兑现的持续参与')
p('会员权益优先使用表情、纪念徽章、月度幕后内容和明确排期的活动，围绕人物关系与共同事件展开。普通观众持续获得完整公开节目。大航海、订阅或其他会员形式以各平台实际可开通能力为准，两平台各自结算，账号名称相同也不能直接当作同一付费用户。')
p('约10万主阵地关注为团队原设的商业化升级节点，不是平台统一准入门槛。需求与交付条件具备时可先做小规模测试；若增长不足，会员和商品收入预测同步顺延。重点观察次月续费、权益使用和运营工时，避免推出过多难以兑现的承诺。')
h('周边从可识别的名场面开始')
p('首批考虑立牌、徽章和纪念卡，将观众确实记住的台词、开播节点或共同任务变为图案。原画、模型、声音和联动素材先确认使用范围，导演核对梗的上下文，运营收集需求后小批量预售。商品写明款式、数量、预计交期与售后，收到回款后仍预留生产和退款资金。')
h('分别核算复购与商品贡献')
p('会员收入按平台净到账减权益直接成本，周边按销售减制作、包装、物流和售后成本计算贡献。全年会员按385个会员月、每月净到账50元测算，50元是情景假设，不是当前平台定价或承诺收益。周边按15,000元销售、10,500元制作履约安排，约4,500元差额还未覆盖团队日常成本。')
h('避免收入与人气重复计算')
p('礼物与付费留言、会员订阅、商品销售、品牌合作分列台账。纪念商品和会员内容可能来自同一场直播，但实际到账只能按订单计一次。每月同时报告续费、退款、履约和可用现金，以持续支持检验角色关系能否形成经营基础。')

# 15 commercial
start('商业模式 品牌直播与游戏合作')
h('品牌主题直播')
p('把合作主题转成角色可以参与的任务。真人介绍产品事实、规则和体验背景，主动设置可展开的挑战；AI以角色视角回应和选择，观众投票或预测。初期拟交付一个完整直播段落、两条重点切片与一份效果报告，具体数量、时长和使用范围在单次合作中确认。')
h('游戏体验合作')
p('围绕已验收的游戏类型组织人机体验回，展示真实选择、讨论和结果。真人解释规则并串场，AI的行动受实际接口与版本能力限制。发布前核对游戏画面、音乐与素材用途，交付直播录屏片段、反应切片和活动总结；未完成适配的游戏先做测试，不能直接承诺可播。')
h('系列栏目和IP授权')
p('当单次合作的观看与回访成立，可扩展为多场连载，让未完成任务带来下次期待。角色、台词或共同事件的授权单独约定素材、用途、期限及衍生范围。直播执行费与素材授权费分别定价，避免一次报价包含无限使用。现阶段先整理样片与权利清单。')
h('从需求到结案的交付流程')
p('市场商务确认目标人群与评价方式，运营组织主持、人格与导演编排，技术确认版本和接管；播前核对品牌事实，播中标记交付时间戳，播后剪辑并提交报告。AI即兴不替代产品信息核验，涉及功效、价格或合作承诺的表达由真人确认。')
h('报价与回款')
p('基准计划仍按首年两次、每次5,000元的目标测算，实际收入以签约及结算为准。报价覆盖策划、主持、模型语音、剪辑、运营和约定授权，逐项核算直接成本与回款周期。先向品牌展示实际受众和节目样例，再谈观看、互动或点击目标；只有取得可归因数据时才报告转化。')

# 16 finance
start('商业模式 预算与首年收支测算')
p('测算覆盖首个完整运营年度，保留原预算口径；双平台费用先在同一预算内试跑，用实际带宽、工具和人力需求复核后再扩大。下表是目标情景，不是已实现收入。')
table(['项目 元','第一季度','第二季度','第三季度','第四季度'],[
    ['礼物与付费留言净到账','1,200','3,000','5,300','7,900'],
    ['会员净到账','0','0','5,250','14,000'],
    ['周边销售','0','0','5,000','10,000'],
    ['商业合作收款','0','0','0','10,000'],
    ['收入合计','1,200','3,000','15,550','41,900'],
    ['日常运营支出','13,000','13,000','13,000','13,000'],
    ['周边制作履约','0','0','3,500','7,000'],
    ['运营现金余额','-11,800','-10,000','-950','21,900'],
],[4.2,3.1,3.1,3.1,3.1])
h('全年现金与资金用途')
p('全年收入61,650元，其中普通打赏17,400元、会员19,250元、周边15,000元、商业合作10,000元。运营支出52,000元，周边制作履约10,500元，运营现金余额为负850元；另计首播资产与工具25,000元，首年现金缺口为25,850元。')
p('10万元启动支持拟用于原画8,000元、Live2D及表情14,000元、开发工具3,000元、年度日常运营52,000元、周边周转10,500元和机动资金12,500元。日常运营含模型语音云服务18,000元、剪辑社群补贴12,000元、推广18,000元及活动杂费4,000元。')
h('投入效率与增长不足时的安排')
p('约250场、500小时排期对应日常现金约208元每场、104元每小时；这不含成员完整薪酬，不能当作完整制作成本。每场另外登记真人主持、策划、开发和剪辑工时。单次合作贡献按净收入减直接成本评估，账号礼物面额不参与现金计算。')
p('按季度并将25,000元启动投入置于期初，基准累计现金缺口峰值约47,750元。若会员与合作延后，按8,000元打赏、40,000元日常支出及25,000元启动投入安排，资金缺口约57,000元；相应推迟预售和可变推广，避免把未到账收入用于长期承诺。')

# 17 dual platform
start('六 市场战略规划 双平台直播')
p('国内运营以B站与抖音双平台直播为目标，共享角色和任务事实，分别设计画面、消息接入与收入记录。初期拟每周二至周六北京时间20:00至22:00开播；周一用于研发和企划，周日复盘与素材整理。')
fig('07-dual-platform','双平台运营方案  同时开播须完成账号权限 推流 音视频与事件联调')
h('两端都是直播阵地')
p('B站用较完整的游戏过程、角色关系和社群讨论承接追更，抖音用快速说明情境、明确反应和阶段结果降低陌生观众的理解成本。抖音包含原生直播与弹幕互动，短片用于直播前后触达。两端分别使用适合平台的横竖画面，不能只把横屏界面机械裁成竖屏。')
h('开播前完成可检验的连接')
p('先核对两个账号的开播权限与适用要求，使用平台允许的推流方式。实际测试音画同步、角色字幕、重复音频、断线恢复及两端观众消息，确认舞台控制不被播出。当前工程的OBS和B站适配是基础，尚不等于已经验收双路持续直播；上线目标和已完成状态分开记录。')
h('统一节目与分别复盘')
p('主持根据标记来源回应两端观众，导演避免重复任务和冲突指令。每周发布4至5条重点切片，保留真人抛梗、AI回应和结局。播放、关注、支持人数与收入分别统计，跨平台账号未经授权匹配不强行合并，以各自回访判断是否值得继续投入。')

# 18 timeline
start('市场战略规划 阶段目标与增长')
h('首月完成可以评价的节目')
p('先完善原创形象与授权声线首版，挑选一类短游戏任务和两个人机栏目，形成3场内部试播及至少10个切片候选。完成中文接梗场景检查和首轮受众体验，逐场保留不成功的触发点、退出原因及技术异常。双平台联调与内容试播并行，但开始固定公开运营以实际验收为前提。')
h('公开运营前三个月建立复播习惯')
p('按已落实人力稳定开播，优先保留观众能记住、愿意接续的栏目，再小额推广表现较好的切片。原计划争取2.5万主阵地关注，执行时同时看七天回访、直播停留与有效关注成本；曝光增长而回访不增长时，先调整内容，再增加投流。')
h('第四至六个月验证可重复的增长')
p('逐步形成游戏主线、挑战、弹幕预测和旧事回收的固定节目结构，选择同量级真人主播联动，比较人机互动与人与人联动怎样共同带来观众。原计划约六个月争取10万主阵地关注；若实际落后，调整商业收入起点和支出，不把目标写成确定结果。')
h('第七至十二个月形成经营记录')
p('在需求成立后测试会员与小批量预售周边，推出有实际受众样例的品牌主题直播。年末15万关注作为进取目标；比规模更需要证明的是会员续费、商品履约、单次合作贡献及再次合作。海外传播放在国内双平台稳定之后，先用字幕样片观察需求，再决定是否投入海外直播。')
h('推广预算跟随内容表现')
p('全年推广预算18,000元分批使用。先比较自然流量下的完播、真实评论与关注，投放后再看回访和直播参与，不能只用曝光汇报成绩。运营按月发布同一口径的观看、支出、回款和改进记录，CEO依据现金与节目效果决定后续投入。')

# 19 feasibility
start('市场战略规划 可行性与验收')
h('工程证据证明哪些部分可行')
p('Minecraft实机记录支持模型规划、动作执行及原木和木剑入库；STS2的2026年10月20日快照记录18次成功动作与一条口播结束回执；Flappy Bird有独立策略验收。上述证据支持不同模块和局部任务可运行，不证明已经能任意游戏长期自主直播。[I1至I3]')
p('首页、游戏工作台与导演工作台的源文件支持界面和调度设计已有基础，记忆候选机制支持来源审核与接续。将它们组合成自然的人机节目仍需现场测试。技术栈的可行性、观众愿看和商业能收入分别验收，避免用其中一项替代另外两项。[I4至I6]')
h('观众验收与接梗复测')
p('首轮30名成年观众按观看习惯分层，随机平衡两段各6分钟样片的观看顺序，规定体验后可自由继续观看或结束。预设迭代门槛仍为自由观看中位数3分钟、至少15人记住角色特点或事件、至少6人在七天内主动回看满60秒。这些是团队决策目标，报告保留有效分母、退出、缺失和技术故障。')
p('目前P001仅形成一份熟人口述前测记录，成年条件、同意和样片体验尚待确认，不计为正式验证结果。后续增加校外陌生观众，对中文接梗和角色关系单独复测，只有达到重复观看与稳定制作要求时才扩大排期和预算。[I7]')
h('把主要失败条件放进执行安排')
p('接梗生硬时回到具体场景与样本，不靠堆流行语修补；语音延迟时缩短回复、调整调度并由真人续场；游戏卡住时停止该任务、核对状态并切到备用栏目。双平台事件不可用时先完成连接问题，避免把本地注入模拟成真实弹幕。')
p('品牌交付有明确范围、素材授权和结算记录；观众画面标明AI角色与真人分工。公司化和全职投入在固定开播、留存、回款及成员职责形成连续记录后再扩大，以实际证据决定下一阶段规模。')

# 20 public refs
start('公开资料与链接引用')
p('公开资料查询日期为2026年10月24日。正文中的事实、团队判断与目标测算分别表述；排期、平台能力和账号权限在执行时再次核对。链接指向具体资料，不使用搜索页。')
refs=[
('[S1] YouTube Culture and Trends  2025年4月17日  虚拟创作者观看样本','https://blog.youtube/culture-and-trends/youtube-culture-trends-report-virtual-creators/','用于说明2024年样本的观看行为，不作AI主播独立市场规模。'),
('[S2] Bilibili  2026年第二季度业绩公告  SEC披露','https://www.sec.gov/Archives/edgar/data/1723690/000119312526371261/d434015dex991.htm','引用月活跃用户约3.71亿，仅作平台规模背景。'),
('[S3] Wei与Tyson  Virtual Stars Real Fans  WWW 2025','https://arxiv.org/abs/2502.01553','研究B站VUP观众与粉丝社区，不能替代本项目受众验证。'),
('[S4] Vedal AI  Neuro sama与Evil Neuro官方站','https://vedal.ai/','用于人物、直播、音乐与商品入口的公开对标。'),
('[S5] Neuro第三方排期样本  Sync2Cal','https://www.sync2cal.com/gaming/twitch-streamers/neuro-sama','样本18:00及19:00 UTC换算为北京时间次日02:00及03:00；页面显示时区可能变化。'),
('[S5补充] Vedal官方Twitch排期入口','https://www.twitch.tv/vedal987/schedule','实际场次与临时变动以官方公告为准。'),
]
for label,url,note in refs:link(label,url);p(note)

# 21 implementation refs
start('平台资料与本项目证据')
refs=[
('[S6] Lumi Nox公开仓库与项目说明','https://github.com/MIO-456/Lumi_Nox','支持其双AI语音、发言调度、记忆与游戏基础的对标。'),
('[S7] Bilibili开放平台直播事件文档','https://openhome.bilibili.com/doc/4/c7b549c9-d5d6-0331-b510-465ed4107e01','用于核对弹幕、礼物与付费事件字段；不是收入结算比例依据。'),
('[S8] 抖音开放平台直播SDK功能介绍','https://open.douyin.com/platform/resource/docs/ability/douyin-live-sdk/introduction','说明礼物等互动能力，不代表本项目已取得接入权限或主播结算条件。'),
('[S9] 抖音直播行为规范','https://sf3-cdn-tos.douyinstatic.com/obj/ies-hotsoon-draft/webcast_all_anchor/webcast_ban_protocol.html','用于直播行为与打赏设计边界；执行以账号适用的现行规则为准。'),
]
for label,url,note in refs:link(label,url);p(note)
h('本项目工程索引')
p('[I1] docs/minecraft-deepseek.md与logs/minecraft-deepseek-live-run.json：规划、真实动作及库存变化。\n[I2] output/sts2/streamer-acceptance.json：2026年10月20日动作与口播回执快照。\n[I3] docs/demo-summary-20261018.md及对应游戏验收文件：Flappy Bird条件与未完成范围。\n[I4] public/index.html、stream.css、games.html、studio.html：舞台、叠加层、游戏与导演界面。\n[I5] src/live/runtime.js与src/brain/capabilities.js：B站事件归一化及平台接入边界。\n[I6] src/brain/memory-candidates.js与相关运行时：记忆来源、人工确认及角色状态。\n[I7] output/kunshan-20261020/research/P001-pretest.json：熟人口述前测及缺失字段。')

# 22 Opus handoff structure, no slide authoring
start('后续演示提炼提纲')
p('以下为商业计划书向演示材料的提炼安排，供后续制作使用。正文按约9页组织，产品介绍4页，其余五部分各尽量1页；封面和结束页如需增加，另行安排。演示采用白底、无动画、文字少且重点突出的形式。')
h('第1页 项目背景')
p('讲清面向谁、提供什么节目、为什么真人与AI需要搭档。以主动造梗、即兴接梗、真实结果和观众回访的流程图为主，对应计划书第3页。')
h('第2页 市场规模与竞品分析')
p('保留少量可引用数据，结合Neuro的内容优势、国内观看时段和中文本土化切口。承认成熟度差距，避免写成Neuro不能说中文，对应第4至5页。')
h('第3至6页 产品介绍')
p('第3页呈现观众体验与人机关系；第4页区分观众界面和后台工作台；第5页用技术到实践流程讲实时聊天、形象、弹幕与游戏；第6页用证据说明已实现能力、中文接梗试验和待开发部分。优先使用界面示意、流程和真实证据，对应第6至9页。')
h('第7页 团队介绍')
p('使用组织图展示CEO、技术总监、运营总监及其余成员岗位，每人说明一个主要交付，对应第10页。')
h('第8页 商业模式')
p('用收入结构图展示礼物与会员、周边、品牌直播和游戏合作；配一个具体互动任务例子，说明游戏为何有流量潜力与节目效果。保留关键收入和成本口径，细节由计划书第11至16页支持。')
h('第9页 市场战略规划')
p('展示B站与抖音双平台直播、切片回流及阶段验收，用时间线说明从内部试播到公开运营和商业合作的推进，对应第17至19页。底部保留关键来源链接，图中注明方案、原型或已验收状态。')

for pp in doc.paragraphs:
    if '\t' in pp.text:
        pp.paragraph_format.tab_stops.clear_all()
        pp.paragraph_format.tab_stops.add_tab_stop(Cm(16.6),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        pp.paragraph_format.space_before=Pt(0)
        pp.paragraph_format.space_after=Pt(6)
        pp.paragraph_format.line_spacing=1.1
        snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');pp._p.get_or_add_pPr().append(snap)

doc.core_properties.title='潮生智能真人与AI共同直播商业计划书'
doc.core_properties.subject='项目背景 市场规模 产品介绍 团队介绍 商业模式 市场战略规划'
doc.core_properties.author='潮生智能团队'
doc.save(OUTPUT)
(QA/'v2-page-plan.json').write_text(json.dumps([{'page':i+1,'title':t} for i,t in enumerate(pages)],ensure_ascii=False,indent=2),encoding='utf-8')
print(OUTPUT)
print('Planned pages:',len(pages),'Tables:',len(doc.tables),'Figures:',len(doc.inline_shapes))
