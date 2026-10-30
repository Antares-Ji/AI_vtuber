from pathlib import Path
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'final' / '昆山杯-潮生智能-商业计划书.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2.1), Cm(1.9)
sec.left_margin = sec.right_margin = Cm(2.15)
sec.header_distance, sec.footer_distance = Cm(.85), Cm(.8)

def font(style, size, bold=False):
    s = doc.styles[style]
    s.font.name = 'Microsoft YaHei'
    s.font.size = Pt(size)
    s.font.bold = bold
    s.font.italic = False
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
    s.paragraph_format.space_after = Pt(5)
    s.paragraph_format.line_spacing = 1.15
font('Normal', 10.5)
font('Title', 28, True)
font('Subtitle', 13)
font('Heading 1', 17, True)
font('Heading 2', 12, True)
font('Heading 3', 10.5, True)
for st in doc.styles:
    if st.type == 1 and st.element.pPr is not None:
        for border in list(st.element.pPr.findall(qn('w:pBdr'))):
            st.element.pPr.remove(border)
for name in ['Heading 1', 'Heading 2', 'Heading 3']:
    doc.styles[name].paragraph_format.keep_with_next = True
doc.styles['Heading 1'].paragraph_format.space_after = Pt(14)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(9)

header = sec.header.paragraphs[0]
header.text = '潮生智能    姬岁歌受众调研与内部试播'
header.style = doc.styles['Normal']
for r in header.runs: r.font.size = Pt(8)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run('2026年10月21日    '); r.font.size = Pt(8)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)

def p(text, bold=False, style=None):
    para = doc.add_paragraph(style=style)
    run = para.add_run(text); run.bold = bold
    para.paragraph_format.widow_control = True
    return para

def h(text):
    doc.add_heading(text, 2)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, w in zip(t.columns, widths): c.width = Cm(w)
    for i, text in enumerate(headers): t.rows[0].cells[i].text = text
    for row in rows:
        for c, text in zip(t.add_row().cells, row): c.text = str(text)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'5'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
    t._tbl.tblPr.append(borders)
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit'); trpr.append(no_split)
        if ri == 0:
            rep=OxmlElement('w:tblHeader'); trpr.append(rep)
        for ci, cell in enumerate(row.cells):
            cell.width = Cm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'), '153E49' if ri==0 else ('F2F6F7' if ri%2==0 else 'FFFFFF')); tcpr.append(shade)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:w'),'80'); e.set(qn('w:type'),'dxa'); margins.append(e)
            tcpr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(1)
                para.paragraph_format.space_before=Pt(1)
                para.paragraph_format.line_spacing=1.12
                if ci==0 and widths[ci]<3: para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.name='Microsoft YaHei'; run.font.size=Pt(9)
                    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
                    run.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '172B34')
                    run.bold = ri == 0
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

chapters=[]
def chapter(title):
    if chapters: doc.add_page_break()
    chapters.append(title)
    doc.add_heading(title, 1)


p('姬岁歌受众调研与内部试播执行手册',style='Title')
p('首轮30人研究方案 问卷 观察表与结果模板',style='Subtitle')
p('执行负责人  张涵 刘瑞麒    统筹  姬祥    技术与素材  项目研发及剪辑工作线')
h('本轮要回答的问题')
p('观众是否愿意把姬岁歌当作持续观看的节目，能否记住角色和共同经历，以及连续经历增强功能是否改善观看选择。研究从实际观看和回访判断需求，再用评分与访谈解释原因。')
table(['设计项','执行口径'],[
['样本','计划30名成年参与者，三层各10名：近期V圈观众、游戏与AI兴趣者、较少看VUP者。团队成员单独做工程验收。'],
['招募','校内不同院系和社群招募，非项目熟人至少20名。记录招募来源；每层按符合条件者报名顺序入组。'],
['补贴','每人固定20元完成补贴，预算600元；提前说明结束、负面评价与自由观看选择均不影响已约定补贴。'],
['时长','前测约3分钟，规定样片共12分钟，后测约4分钟，自由观看区间最多10分钟，七天后简短回访。'],
['版本','A为基本互动版，B为连续经历增强版；保持主持人、任务类型、声音、画面质量与长度尽量一致。'],
['顺序','每层5人A后B、5人B后A，总体15人AB、15人BA；使用中性样片代号。'],
['阶段状态','方案已完成；正式招募、观看、问卷和七天回访待执行。']
],[3.0,13.7])
h('准备完成后再开始招募')
p('冻结两个样片、版本编号和播放链接，完成音量与清晰度检查。准备无广告的本地播放或受控入口、匿名编号表和计时记录。先由团队工程人员试跑流程，修复记录问题，再开启正式观众测试。')
p('研究介绍统一说明这是一档由AI角色与真人搭档共同完成的节目。正式参与者自愿加入，研究结果以汇总形式使用，匿名原话引用另获同意。联系方式单独保存，完成回访30天后清理；研究记录使用编号。')
doc.add_page_break()
doc.add_heading('一 招募筛选与前测问卷',1)
p('开场说明：我们正在研究一档游戏互动节目，想了解真实观看体验。你可以随时结束，也可以直接给出负面意见。我们会记录观看选择与匿名反馈；录像、联系方式和原话引用分别征求同意。')
table(['题号','问题与回答方式'],[
['Q1','是否年满18岁并同意参与本次匿名观看研究？是／否。选择否则结束。'],
['Q2','过去30天，你看VUP直播或切片的频率？几乎没有／少于每周1次／每周1至3次／每周4次以上。'],
['Q3','过去30天主要看过哪些内容？游戏直播／VUP杂谈／音乐或翻唱／AI内容／其他，可多选。'],
['Q4','通常一次看直播或长视频多长时间？少于5分钟／5至15分钟／16至30分钟／超过30分钟。'],
['Q5','请写一个你最近主动回看的创作者，以及你回看的具体原因。开放题。'],
['Q6','你对AI虚拟主播的熟悉程度？从未看过／看过片段／偶尔看／持续关注。是否看过Neuro-sama或Lumi/Nox？分别记录。'],
['Q7','你是否认识本项目成员，或参与过项目制作？不认识／认识但未参与／参与过。参与制作转入工程测试组。'],
['Q8','过去30天是否为创作者付费？没有／礼物打赏／月度会员／周边／其他，可多选；可跳过。'],
['Q9','是否同意研究用计时记录？是／否。是否同意联系一次七天回访？是／否。联系方式另表保存。']
],[2.1,14.6])
h('分层规则')
p('优先将每周至少看一次VUP内容者列入V圈层；其余有游戏或AI兴趣者进入游戏AI层；剩余参与者列入低接触层。每人进入一个层级。项目熟人最多10名，结果同时汇总熟人与陌生观众。')
h('招募记录')
table(['编号','受众层','招募来源','熟人关系','观看顺序'],[['','','','',''],['','','','','']],[2.2,3.5,4.2,3.4,3.4])
doc.add_page_break()
doc.add_heading('二 样片设计与现场流程',1)
h('内部版本比较')
p('两个6分钟样片围绕同类短游戏任务。A展示基本聊天与当场行动；B在相同任务结构上接续已确认的共同事件或约定。两版保持主持、剪辑密度、声音与画面质量一致。差异材料和版本编号记录在表中，使用隔离的会话配置，避免跨版本记忆污染。')
p('人工准备的剧情提示标明为研究样片素材；真实执行部分保留工具和状态日志。若增强功能尚未完成，先补齐样片与工程验证，再开展对应功能测试。')
table(['环节','研究人员操作','记录'],[
['开始前','说明AI与真人分工，完成前测和匿名编号。','同意、受众层、熟悉程度、顺序。'],
['样片一','播放6分钟，可随时结束；研究人员保持安静。','开始结束、实际观看、退出与技术问题。'],
['短反馈','先自由回忆，再回答评分题。','记住的事件、角色印象、评分。'],
['样片二','播放另一版本，使用相同步骤。','同样口径记录。'],
['自由选择','提供继续看姬岁歌、其他同类内容或结束三个选择；可随时切换。','10分钟内真实选择和姬岁歌前台观看时长。'],
['结束与回访入口','完成比较题，提供自愿回看链接与可选回访安排。','链接编号、同意联系、补贴完成情况。']
],[3.0,8.0,5.7])
h('普通VUP与竞品作为另一项内容观察')
p('准备相近游戏或杂谈类型、相近长度与音量的公开片段，记录观众原先是否熟悉。先确认素材使用范围，内部播放保留来源。该观察收集观看偏好与原因，单独归档；内部A/B版本比较专门检验增强功能。')
h('样片冻结表')
table(['版本代号','文件或链接','长度','变量与制作说明'],[['A','','6分钟',''],['B','','6分钟','']],[2.5,6.3,2.5,5.4])
doc.add_page_break()
doc.add_heading('三 观看后问卷',1)
p('每个版本播放后独立填写。先回答开放题，再看评分；研究人员按原话记录。评分为1至5：很不同意、不同意、一般、同意、很同意；另设无法判断。')
table(['题号','问题'],[
['P1','刚才发生了什么？请用自己的话描述一个记住的片段。'],
['P2','你觉得这个角色有哪些特点？请举出刚才的具体表现。'],
['P3','哪一刻最有趣，哪一刻最想离开？能回忆大致时间吗？'],
['P4','我能够理解角色正在做什么、为什么这样做。1至5或无法判断。'],
['P5','角色的说话与前后行为让我觉得一致。1至5或无法判断。'],
['P6','真人和AI之间的互动值得继续看。1至5或无法判断。'],
['P7','我能看见观众建议或此前事件对节目产生影响。1至5或无法判断。'],
['P8','声音、节奏和画面让我容易跟上节目。1至5或无法判断。'],
['P9','如果现在有新的同类片段，我愿意继续观看。1至5或无法判断。'],
['P10','有哪一句话、一次反应或一次动作让你觉得不自然、重复或难理解？开放题。']
],[2.2,14.5])
h('两版比较题')
p('C1：你更愿意继续看A、B、两者差不多，还是都想结束？为什么？C2：你觉得两版最大的区别是什么？C3：如果将来每周固定开播，你会因为什么回来，又会因为什么不来？')
h('商业需求放在内容评价之后')
p('C4：达到你愿意支持的程度，节目还需要改善什么？C5：如果未来提供月度幕后内容、表情、纪念周边，你最看重什么？可以选择都不需要。C6：愿意考虑的支持方式与金额区间，可跳过。意愿题与观看、回访及真实付费分别保存。')
p('自由观看结束后追加：你刚才为什么做出这个选择？如果继续看，最想知道什么；如果换内容或结束，是什么原因？')
doc.add_page_break()
doc.add_heading('四 观察记录与数据字段',1)
p('使用同一匿名编号关联问卷、样片和回访。原始记录保留开始结束时间与时区，按ISO 8601格式保存；前台播放时长与规定观看时长分开记录。')
table(['字段','填写规则'],[
['participant_id cohort source','匿名编号、受众层、招募来源与熟人关系；原始联系方式放独立表。'],
['version order session_id','样片版本、AB或BA顺序、录屏或日志场次。'],
['required_watch_sec','规定样片实际观看秒数；提前退出按实际记录。'],
['optional_watch_sec','自由区间内姬岁歌前台有效播放秒数，累计上限600；选择结束记0，缺失留空。'],
['choice reason timestamp','继续、换内容或结束；按原话记录原因和对应时刻。'],
['recall_text recall_code','原话保留；两名标注者独立判断是否准确回忆具体事件或角色表现，分歧再复核。'],
['score_P4_to_P9','原始1至5评分；无法判断单列，保留每题有效人数。'],
['tech_issue intervention','网络、语音、画面、游戏或现场干预，记录起止与影响。'],
['return7d watch7d_sec','七天窗内自愿回看及累计时长；未同意可识别记录者记缺失。'],
['quote_permission artifact_ref','匿名原话引用授权与对应文件、录屏时间戳或日志引用。']
],[5.4,11.3])
h('现场空白观察表')
table(['编号与版本','时刻','观察到的行为','技术情况','原话与记录者'],[['','','','',''],['','','','',''],['','','','','']],[3.0,2.0,4.5,3.0,4.2])
h('记录完整性')
p('每次结束即核对匿名编号、版本和时间戳。观看不满规定样片者仍保留退出资料，结果分列完整体验与提前退出。技术故障单列，主报告说明有效样本、退出与缺失人数，并保留包含所有入组者的敏感性对照。')
doc.add_page_break()
doc.add_heading('五 指标计算与复盘门槛',1)
table(['指标','分子或计算方式','分母与注意事项'],[
['自由观看中位时长','自由区间姬岁歌有效播放秒数中位数。','有选择机会且计时有效的参与者；选择结束为0，技术缺失单列。'],
['角色或事件记忆率','准确回忆至少一项具体事件或角色表现的人数。','完成相应提问的有效人数；写明双人复核规则。'],
['继续选择比例','自由选择首先选姬岁歌的人数。','获得相同选择机会的有效人数。'],
['B偏好比例','明确更愿意继续看B的人数。','有两版有效体验与比较答案者；平手和都不喜欢分别展示。'],
['七天实际回访率','七天内主动回看姬岁歌累计至少60秒的人数。','同意回访计时且入口可用者；另报其占全部入组人数。'],
['评分分布','各题1至5的分布、中位数和有效人数。','无法判断与缺失分列。按受众层和熟人关系分组。']
],[4.1,6.0,6.6])
h('预先冻结首轮决策目标')
p('计划至少30份有效体验；自由观看中位时长达到180秒；至少15人准确回忆具体角色或事件；至少6人在七天内主动回看满60秒。增强功能同时观察B偏好是否达到18人，以及受众说出的改善是否与设计变量对应。门槛是团队首轮筛选目标，后续依据实际分布和复测完善。')
p('达标时扩展稳定节目，并补充陌生校外受众；观看短、回访少时优先改节目与定位。若只有某层表现好，先扩大该层验证。技术故障明显影响体验时，先处理故障再用新参与者复测。每轮只改变清楚记录的一组因素。')
h('呈现结果的方法')
p('先报告人数和原始分布，再给比例；分别呈现正面与负面反馈。便利样本说明招募范围，组间差异作为方向线索。通过匿名原话、录屏片段和回访行为共同解释结果。补贴、邀请和回访提醒的发送时点均保留记录。')
doc.add_page_break()
doc.add_heading('六 七天回访与深度访谈',1)
h('七天观察窗')
p('D0结束后提供相同的自愿回看入口。D0至D7记录主动回看，区分通知或提醒后的访问。D7观察窗关闭后，再向同意联系者开展简短回访。公开展示使用匿名汇总；个人记录按参与说明管理。')
table(['问题','要了解的事'],[
['这一周你是否又看过姬岁歌？从哪里点进去？','核对实际入口与记录，区分主动回访和提醒访问。'],
['你现在最记得哪件事或哪句话？','观察一周后仍留下的角色与节目记忆。'],
['如果又看了，是想看什么后续？','识别共同经历、任务、人物关系等回访动机。'],
['如果没看，是什么原因？','时间、入口、内容兴趣和节目质量分别记录。'],
['你平常会用哪个创作者的内容替代这段时间？','了解真实注意力竞争。'],
['与熟悉的VUP或AI主播比，哪一部分更想看？','收集具体差异和已有观看经验。'],
['哪一部分你觉得换个普通聊天模型也可以做到？','识别真正需要产品支持的价值。'],
['你会怎样向朋友介绍姬岁歌？','检验角色和产品记忆点是否清楚。'],
['你希望下一场出现什么，也希望删掉什么？','将需求转化为可制作的栏目。'],
['什么改变会让你愿意持续看，或考虑支持？','理解持续观看与商业支持之间的条件。']
],[8.6,8.1])
h('访谈人数与整理')
p('计划从完成体验者中选择10人深访，覆盖继续看与退出、回访与未回访、熟人与陌生观众。按原话整理主题，每个主题附人数、受众层和片段时间戳。保留与团队预期相反的意见，列入下一轮问题。')
doc.add_page_break()
doc.add_heading('七 结果报告模板与竞品采样',1)
p('研究状态  待招募。以下字段由真实记录填写；空白表示尚待采集。')
table(['结果字段','实际数据','证据位置'],[
['招募人数 有效人数 退出与缺失','',''],
['受众分层 熟人和陌生人数','',''],
['自由观看中位数与分布','',''],
['具体事件与角色回忆人数','',''],
['A偏好 B偏好 平手 都不喜欢','',''],
['七天实际回访人数与分母','',''],
['代表性正面反馈与授权','',''],
['代表性负面反馈与授权','',''],
['下一轮修改与责任人','','']
],[7.0,4.0,5.7])
h('普通VUP与AI主播采样表')
p('为hanser、小缘、泠鸢、兰音、永雏塔菲、雫るる、MeUmy、宝钟马琳、Neuro-sama及Lumi/Nox分别建表。页面截取标明检索日、缓存或发布日期、平台和账号；组织关系留存公开声明，公会合同采用公开披露或直接访谈信息。')
table(['字段','统一规则'],[
['账号与组织','个人或小团队主导、公开机构合作、公会服务、企业所属分列。'],
['关注与作品规模','保留页面快照和时间；组合的多个账号分别记录。'],
['近30天内容','最近10条符合条件的官方短片，音乐、切片、公告分列；少于10条全部保留。'],
['普通表现','记录每条播放、点赞、评论、时长和发布天数，计算中位数。'],
['90天节奏','统计各类型更新量；直播、录像和短片分开，记录重大活动。'],
['内容学习','互动方式、回访理由、可复制的制作方法及本项目的下一项验证。']
],[4.0,12.7])
p('向评委展示时，先呈现外部市场行为和项目功能证据，再汇报实际测试进度、方法和结果。每个优势配一个可看的片段、一个可核对的记录和一项观众行为。')

OUT=ROOT/'final'/'姬岁歌-受众调研与内部试播执行手册.docx'
doc.core_properties.title='姬岁歌受众调研与内部试播执行手册'
doc.core_properties.subject='首轮30人研究方案与问卷记录模板'
doc.core_properties.author='潮生智能团队'
doc.save(OUT)
print(OUT)
