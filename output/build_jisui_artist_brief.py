from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"E:\Github\AI_vtuber")
IMAGE = Path(r"C:\Users\Antares.Ji\.codex\generated_images\01a0e670-ad4c-7c20-9280-655e542412e0\exec-5aee6f94-dd5c-4b36-9668-92168568e099.png")
OUT = ROOT / "docs" / "design" / "姬岁歌-画师交接稿-2026-10-04.docx"


def set_font(run, name="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = name
    if run._element.rPr.rFonts is not None:
        run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
        run._element.rPr.rFonts.set(qn("w:ascii"), name)
        run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in list(tc_pr.findall(qn("w:shd"))):
        tc_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(table):
    tbl_pr = table._tbl.tblPr
    bd = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + side)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "5")
        el.set(qn("w:color"), "D9D9D9")
        bd.append(el)
    tbl_pr.append(bd)


def margins(cell, top=110, start=125, bottom=110, end=125):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        item = OxmlElement("w:" + name)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")
        mar.append(item)
    tc_pr.append(mar)


def para(doc, text="", style=None, before=0, after=5, line=1.28):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    set_font(p.add_run(text), size=10.3)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.23
    set_font(p.add_run(text), size=10.1)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(11 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    set_font(p.add_run(text), size=15 if level == 1 else 11.5, bold=True, color="000000")
    return p


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    borders(t)
    if widths:
        for col, w in zip(t.columns, widths):
            col.width = Cm(w)
    for idx, text in enumerate(headers):
        c = t.rows[0].cells[idx]
        shade(c, "333333")
        margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(text), size=9.3, bold=True, color="FFFFFF")
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for idx, text in enumerate(row):
            c = cells[idx]
            if ridx % 2:
                shade(c, "F7F7F7")
            margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.18
            set_font(p.add_run(text), size=9.1)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(1.7)
sec.bottom_margin = Cm(1.6)
sec.left_margin = Cm(2.15)
sec.right_margin = Cm(2.15)

styles = doc.styles
for sname in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet"):
    style = styles[sname]
    style.font.name = "Microsoft YaHei"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(10.3)
styles["Title"].font.size = Pt(22)
styles["Title"].font.bold = True
title_style_ppr = styles["Title"]._element.get_or_add_pPr()
for old in list(title_style_ppr.findall(qn("w:pBdr"))):
    title_style_ppr.remove(old)

title = doc.add_paragraph(style="Title")
title.paragraph_format.space_after = Pt(8)
set_font(title.add_run("姬岁歌角色视觉设计画师交接稿"), size=22, bold=True, color="000000")

p = para(doc, "AI 虚拟主播角色概念  2026 年 10 月 4 日  画师沟通版", after=9)
for r in p.runs:
    r.font.color.rgb = RGBColor.from_string("555555")

para(doc, "本稿用于把姬岁歌目前确认的视觉方向交给原画师深化。目标是一位有少女感的 AI 虚拟主播：以黑红裙装为主体，融入少量蜜蜡黄色、绢白和旧金；历史细节落在利簋启发的腰扣，科技身份通过发卡、袖口与侧腰的声纹装置呈现。下图是沟通用概念图，文字说明优先于图中的偶发细节。", after=8)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
inline = p.add_run().add_picture(str(IMAGE), width=Cm(14.2))
inline._inline.docPr.set("descr", "姬岁歌单人服饰概念图，展示黑红裙装、黄色裙褶、发卡、异色瞳、青铜腰扣及声纹装置")
p.paragraph_format.space_after = Pt(3)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(0)
set_font(p.add_run("图一  当前视觉沟通稿  AI 生成草图  供结构与元素讨论"), size=8.5, color="666666")

doc.add_page_break()
heading(doc, "角色与整体造型")
para(doc, "姬岁歌的故事设定是同一身体中有“岁”和“歌”两位虚构人格。当前画的是她们共用的一套主视觉服饰，不是把两位角色左右拼接，也不是反派黑化造型。面部气质应年轻、好奇，带一点有主见的小傲娇；避免成熟妩媚或幼态化。")
para(doc, "视觉优先级：先读到少女主播的完整裙装与自然发型，再看到红黑主色和黄色衣料，最后发现青铜、玉与声纹科技细节。衣服要像真实可裁剪的布料，有清楚的领、袖、腰和裙层关系。")

heading(doc, "服饰与色彩")
table(doc, ["部位", "画师执行要点"], [
    ("上身", "黑色斜襟短上衣为主，绢白窄领形成脸部亮区；肩颈完整包覆。袖套有中式宽袖的流动感，内层袖保持操作游戏时的利落。"),
    ("裙身", "必须是裙子。黑红层叠裙摆保持主导，左侧加入一片可辨认的蜜蜡黄色裙褶或内片，约占裙身可视面积的一小部分；红与黄由黑色和绢白过渡，不做大面积纯色拼旗。"),
    ("双拖尾", "从后腰分出两条独立布带，随动作飘动。红色一条更醒目，另一条用烟黑透纱与少量暖黄内衬；形状圆润柔软，不像燕尾服，也不似旗帜。"),
    ("腿部与鞋", "黑色不透明袜或丝袜，腿饰佩戴在袜面；低跟短靴适合直播与游戏动作。鞋上少量旧金五金件与红色收边即可。"),
], widths=[3.0, 13.7])

palette = table(doc, ["色彩", "参考色", "用途"], [
    ("玄墨黑", "#252024", "上衣、裙层、袜与鞋的结构底色"),
    ("朱漆绛红", "#A23F45", "裙摆、袖内与主拖尾；有漆色深度，避免纯正大红"),
    ("蜜蜡黄", "#C99A4A", "左侧裙褶、少量拖尾内衬；应是布料颜色，不只金属镶边"),
    ("暖绢白", "#F3EBDD", "斜襟领边与小面积内衬"),
    ("旧铜金", "#9B7749", "腰扣、发卡与小五金；不等同于黄色布料"),
], widths=[3.2, 3.1, 10.4])
for row_idx, fill in enumerate(("252024", "A23F45", "C99A4A", "F3EBDD", "9B7749"), start=1):
    cell = palette.rows[row_idx].cells[0]
    shade(cell, fill)
    for run in cell.paragraphs[0].runs:
        run.font.color.rgb = RGBColor.from_string("000000" if fill in ("C99A4A", "F3EBDD") else "FFFFFF")
para(doc, "色值是沟通范围，不是最终印刷或建模标准。请画师在保持黑红主次与黄色可见度的前提下调整材质、明度和肤色关系。", after=0)

doc.add_page_break()
heading(doc, "头发与面部")
bullet(doc, "头发为自然的深色半扎长发，偏分刘海；发量和发束疏密有变化，避免左右镜像的大卷、重复细辫与密集发光线。刘海只留一束清楚的深红条染。")
bullet(doc, "头上一枚单侧发卡：旧金与浅玉结合，造型借谷穗的弧线。发卡边缘可有小型声谱灯，但不做猫耳、头戴耳机或冠饰。")
bullet(doc, "眼睛保留红与暖金的自然异色瞳，虹膜细节有层次，瞳孔为普通圆形。不要龙瞳、竖瞳、五角星或过度浑圆的眼型。")

heading(doc, "历史元素的落点")
para(doc, "腰部正中是全身唯一的大件历史识别点。参考西周利簋的兽首双耳、方座与器身兽面纹，把双耳曲线和兽面提炼成可穿戴的青铜腰扣；四角的蝉纹可缩成小刻痕。腰扣要像精制饰件，不把整件青铜器挂在身上，也不照描铭文。国博介绍见文末来源。")
para(doc, "侧腰的小玉环称“渭水玉璧”，是角色原创饰物，借用玉璧的圆形与渭水意象，并非真实出土文物。让它与腰扣错位摆放，避免胸前出现大圆盘。")

heading(doc, "AI 虚拟主播元素")
table(doc, ["位置", "表现方式", "视觉边界"], [
    ("发卡", "五格以内细小声谱灯，靠近浅玉粒与金属边缘。", "远看仍是发卡，不像天线或猫耳。"),
    ("单侧袖口", "一片贴合袖口的半透明音频控制片，能看到少量电平与声波。", "不要整只手臂铺满电路或悬浮大屏。"),
    ("侧腰玉环", "玉环旁有一圈局部发光的声纹刻度，形成“历史材质＋数据界面”。", "体积小于掌心，与正中青铜扣分开。"),
    ("拖尾末端", "每条拖尾仅一道细声波线，随布料弯曲。", "不画音符、文字、旗帜与密集光带。"),
], widths=[2.5, 8.2, 6.0])
para(doc, "这些是角色外观设定，可先用静态贴图或轻量动效表达。当前项目尚未完成歌唱模块，不要把概念图解读为已实现的实时歌唱功能。", after=0)

doc.add_page_break()
heading(doc, "原画深化与交付")
para(doc, "请优先把概念图里模糊的结构画清楚，再调整装饰密度。建议先交一版简化线稿和上色稿确认轮廓，再制作完整设定页。")
table(doc, ["建议交付", "需要画清楚的内容"], [
    ("正侧背三视图", "上衣斜襟、袖套连接、裙摆左黄右红分区、后腰两条拖尾的起点与层次。"),
    ("头部放大图", "自然发束、单束刘海条染、单侧发卡、普通瞳孔的红金异色瞳。"),
    ("饰件详图", "利簋启发的腰扣、渭水玉璧、袖口音频片、鞋与腿饰。标注金属、玉、绢与透纱的材质区别。"),
    ("表情与拆件建议", "好奇、忍笑、倔强、专注、生气、疲惫等基础表情；头发、袖套、裙层、双拖尾与声纹灯应便于后续 Live2D 分层。"),
], widths=[3.8, 12.9])

heading(doc, "本轮已排除的方向")
bullet(doc, "不要把岁与歌画成两位并排少女，也不要做简单黑白双生或邪恶化形象；当前主视觉是一人一套融合服饰。")
bullet(doc, "不要猫耳耳机、巨大胸前发光圆环、龙瞳竖瞳、复制感强的发束、成熟露肩造型、满身青铜纹或发光电路。")
bullet(doc, "不要借现有游戏、漫画或虚拟歌姬的具体服饰、符号和五官；炎黄与西周是文化叙事线索，画面不使用旗帜、徽章、星形或政治口号。")

heading(doc, "参考与状态")
para(doc, "文物依据：中国国家博物馆“利”青铜簋馆藏页，https://www.chnmuseum.cn/zp/zpml/kgfjp/202108/t20210802_250931.shtml 。利簋属于西周初年，馆藏说明记载兽首双耳、方座以及器身和方座的兽面纹。")
para(doc, "角色依据：项目中的《姬岁歌 歌人格设定》《姬岁歌 岁人格设定》。本稿是画师沟通版；最终色值、布料裁片、表情与 Live2D 拆件须由原画师和项目组评审确认。图一为 AI 生成的视觉草图，仅用于表达方向，不作为可直接投放的正式原画。", after=0)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(footer.add_run("姬岁歌 画师交接稿  ·  2026 10 04"), size=8, color="777777")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
