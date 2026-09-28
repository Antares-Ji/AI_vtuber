from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = 'AI虚拟主播项目简述与进度.docx'

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.2)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)

normal = doc.styles['Normal']
normal.font.name = 'Microsoft YaHei'
normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.45
normal.paragraph_format.space_after = Pt(5)

for name, size in [('Title', 19), ('Heading 1', 13)]:
    style = doc.styles[name]
    style.font.name = 'Microsoft YaHei'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(11 if name == 'Heading 1' else 0)
    style.paragraph_format.space_after = Pt(5)

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('AI虚拟主播项目简述与当前进度')

author = doc.add_paragraph()
author.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = author.add_run('创作人：姬祥')
run.font.name = 'Microsoft YaHei'
run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
run.font.size = Pt(10.5)
author.paragraph_format.space_after = Pt(12)

sections = [
    ('一 项目简述', '本项目是 Live2D AI 虚拟主播原型，面向直播陪伴和角色化互动。系统整合形象、弹幕、语言模型、语音、情绪人格和记忆，并以工作台管理状态、故事和直播链路。'),
    ('二 建设目标', '目标是形成稳定人设和自然回应：优先级队列处理互动，分层记忆保存确认信息，情绪与导演策略决定语气和表演；本地或兼容模型、GPT SoVITS、SenseVoice 构成语音闭环，并落实隐私最小化、失败降级和人工确认。'),
    ('三 当前进度', 'Live2D 舞台、弹幕、34维人格、63维情绪、SQLite 记忆、故事存档、工作台及 OBS 入口已完成。实时链路已记录首字延迟、语音首字节和播放指标；本地与云端模型可双向降级，门控预取保持播放顺序。核心回归已通过。'),
    ('四 现阶段边界与下一步', '项目仍为原型：原生流式 ASR 因缺少本地权重关闭，浏览器使用批量识别；B站正式连接需平台凭据；双机 Worker 仅支持只读注册与心跳；明日方舟视觉已完成本地 OCR 和主界面样本验证，尚无多场景识别或自动操作。下一步进行真实设备验收、ASR 评估、Worker 上报和多场景验证。'),
]

for heading, body in sections:
    doc.add_paragraph(heading, style='Heading 1')
    p = doc.add_paragraph(body)
    p.paragraph_format.first_line_indent = Cm(0.74)

doc.core_properties.author = '姬祥'
doc.core_properties.title = 'AI虚拟主播项目简述与当前进度'
doc.save(OUT)
