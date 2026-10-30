"""在《昆山杯-潮生智能-路演PPT-1025稿.pptx》基础上增补修改为 16 页（不删除原有页面，原 10 页相对顺序保留）。
用法：py -3 revise_deck.py <原稿.pptx> <输出.pptx>
"""
import sys, copy, re
from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
from lxml import etree

SRC, OUT = sys.argv[1], sys.argv[2]
ASSET_LIVE = r"E:\Github\AI_vtuber\output\kunshan-20261020\ppt\assets\live-focus.png"  # 开发版运行截图 2026-10-14
ASSET_STUDIO = r"E:\Github\AI_vtuber\output\kunshan-20261020\ppt\assets\studio.png"    # 导演工作台记忆审核页截图

ACC = RGBColor(0x24, 0x59, 0xD6); INK = RGBColor(0x1F, 0x23, 0x29); MUTED = RGBColor(0x5B, 0x62, 0x70)
FAINT = RGBColor(0x8A, 0x90, 0x9C); LINE = RGBColor(0xD5, 0xD9, 0xE0); SOFT = RGBColor(0xF4, 0xF6, 0xF9)
ACCSOFT = RGBColor(0xEA, 0xF0, 0xFD); WHITE = RGBColor(0xFF, 0xFF, 0xFF); ACC3 = RGBColor(0xB9, 0xCC, 0xF4)
FONT = "Microsoft YaHei"
TOTAL = 16

prs = Presentation(SRC)
orig = list(prs.slides)  # 原 1–10 页


# ---------- helpers ----------
def font(run, size=16, bold=False, color=INK):
    f = run.font
    f.size = Pt(size); f.bold = bold; f.color.rgb = color; f.name = FONT
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def write(tf, paras, size=16, bold=False, color=INK, align="l", anchor=None, lsp=None, space=None):
    """paras: str（\n 分段）或 [段落]；段落为 str 或 [(text, {size,bold,color,link})]"""
    if isinstance(paras, str):
        paras = paras.split("\n")
    tf.clear()
    tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if lsp: p.line_spacing = lsp
        if space: p.space_after = Pt(space)
        runs = [(para, {})] if isinstance(para, str) else para
        for text, o in runs:
            r = p.add_run(); r.text = text
            font(r, o.get("size", size), o.get("bold", bold), o.get("color", color))
            if o.get("link"):
                r.hyperlink.address = o["link"]


def tb(s, x, y, w, h, paras, size=16, bold=False, color=INK, align="l", anchor="t", lsp=None, space=None, name=None):
    t = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = t.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    write(tf, paras, size, bold, color, align, anchor, lsp, space)
    if name: t.name = name
    return t


def box(s, x, y, w, h, fill=WHITE, line=LINE, lw=1.0, dash=False, rad=0.06, name=None):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rad else MSO_SHAPE.RECTANGLE, I(x), I(y), I(w), I(h))
    if rad: sh.adjustments[0] = rad
    if fill is None: sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
        if dash: sh.line.dash_style = MSO_LINE.DASH
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    if name: sh.name = name
    return sh


def line(s, x1, y1, x2, y2, color=FAINT, w=1.5, arrow=False, dash=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    c.line.color.rgb = color; c.line.width = Pt(w)
    if dash: c.line.dash_style = MSO_LINE.DASH
    if arrow:
        ln = c.line._get_or_add_ln()
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")
    return c


def pill(s, text, x, y, w, kind="done", size=12, h=0.32):
    st = {"done": (ACC, None, WHITE, False), "wip": (WHITE, ACC, ACC, False), "plan": (WHITE, FAINT, MUTED, True),
          "soft": (ACCSOFT, None, ACC, False), "grey": (SOFT, None, MUTED, False)}[kind]
    box(s, x, y, w, h, fill=st[0], line=st[1], lw=1.25, dash=st[3], rad=0.5)
    tb(s, x, y, w, h, text, size=size, bold=True, color=st[2], align="c", anchor="m")


def header(s, label, title, sub=None, tsize=30):
    tb(s, 0.6, 0.32, 9, 0.3, label, size=13, bold=True, color=ACC)
    tb(s, 0.6, 0.62, 12.13, 0.72, title, size=tsize, bold=True, anchor="m")
    if sub: tb(s, 0.6, 1.36, 12.13, 0.4, sub, size=16, color=MUTED)


def footer(s, n, src=None):
    line(s, 0.6, 6.98, 12.73, 6.98, color=LINE, w=0.75)
    if src: tb(s, 0.6, 7.04, 10.8, 0.3, src, size=11, color=FAINT, anchor="m")
    pagenum(s, n, new=True)


def pagenum(s, n, new=False):
    paras = [[(f"{n:02d}", {"bold": True, "color": INK}), (f" / {TOTAL}", {"color": FAINT})]]
    if new:
        tb(s, 11.53, 7.04, 1.2, 0.3, paras, size=12, align="r", anchor="m", name="page-number")
        return
    for sh in s.shapes:
        if sh.has_text_frame and re.fullmatch(r"\s*\d{2}\s*/\s*\d{1,2}\s*", sh.text_frame.text):
            write(sh.text_frame, paras, size=12, align="r", anchor="m")
            sh.left, sh.top, sh.width, sh.height = I(11.53), I(7.04), I(1.2), I(0.3)
            return
    raise RuntimeError("page number not found")


def find(s, sid):
    for sh in s.shapes:
        if sh.shape_id == sid:
            return sh
    raise KeyError(sid)


def move(sh, x=None, y=None, w=None, h=None):
    if x is not None: sh.left = I(x)
    if y is not None: sh.top = I(y)
    if w is not None: sh.width = I(w)
    if h is not None: sh.height = I(h)


def remove(sh):
    sh._element.getparent().remove(sh._element)


def notes(s, text, keep_old=True):
    ns = s.notes_slide
    if ns.notes_text_frame is None:   # 新增页的备注页缺少正文占位符：从原第 1 页备注复制一个
        src = orig[0].notes_slide.notes_placeholder._element
        ns.shapes._spTree.append(copy.deepcopy(src))
    tf = ns.notes_text_frame
    old = tf.text if keep_old else ""
    tf.text = text + (("\n\n【原备注（保留；其中财务测算等历史版本内容按标注理解）】\n" + old) if old.strip() else "")


def new_slide():
    s = prs.slides.add_slide(prs.slide_layouts[0])
    for ph in list(s.placeholders):
        remove(ph)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    return s


# ===================== 原第 1 页 → 新 01 封面 =====================
s = orig[0]
write(find(s, 5).text_frame, [[("姬岁歌", {"size": 88, "bold": True})]], size=88, bold=True)
move(find(s, 5), 1.1, 1.75, 5.2, 1.5)
tb(s, 1.1, 3.35, 5.3, 1.05, "面向中文游戏观众的\n人机共播虚拟角色项目", size=26, bold=True, lsp=1.1)
move(find(s, 6), 1.1, 4.7)                 # 蓝色短线
move(find(s, 2), 1.1, 4.88, 5.2, 0.5)       # 潮生智能
write(find(s, 2).text_frame, "潮生智能", size=22, bold=True, color=MUTED)
tb(s, 1.1, 5.5, 5.3, 0.9, [[("当前阶段  ", {"bold": True, "color": ACC}), ("原型开发；内部试播与观众验证待开展", {})],
                           [("形象  ", {"bold": True, "color": ACC}), ("概念图定稿中；原画与 Live2D 计划 2026-11 底交付", {})]],
   size=13, color=MUTED, space=4)
pic = find(s, 9)
pic.crop_right = 0.46                        # 保留主立绘，细节拼图留在原图中（裁切可在 PPT 里还原）
ih = 6.0; iw = ih * (1226 * 0.54) / 1283
move(pic, 12.73 - iw - 0.3, 0.55, iw, ih)
move(find(s, 10), 12.73 - iw - 0.3, 6.6, iw, 0.3)
write(find(s, 10).text_frame, "角色概念图（定稿中）", size=12, color=FAINT, align="c")
pagenum(s, 1)
notes(s, "【增补修改版 · 讲述重点】\n· 一句话定位：姬岁歌是面向中文游戏观众的人机共播虚拟角色项目——主持人设局造梗，AI 以岁、歌两种性格接梗，围绕游戏与弹幕形成节目。\n· 当前阶段如实说明：原型开发阶段，内部试播与观众验证尚未开展；角色形象定稿中，原画与 Live2D 计划 2026 年 11 月底交付。\n· 右侧为角色概念图主立绘（原图细节拼图通过裁切隐藏，可在 PowerPoint 中重置裁切恢复），不是 Live2D 运行画面。\n· 承接：下一页用一页摘要回答“做什么、给谁看、做到哪、要多少钱换什么”。")

# ===================== 新 02 项目摘要 =====================
s2 = new_slide()
header(s2, "项目摘要", "姬岁歌的节目定位与首阶段验证")
cw, ch, x0, y0 = 5.92, 2.42, 0.6, 1.62
cards = [(x0, y0), (x0 + cw + 0.29, y0), (x0, y0 + ch + 0.25), (x0 + cw + 0.29, y0 + ch + 0.25)]
# 1 做什么
x, y = cards[0]; box(s2, x, y, cw, ch)
tb(s2, x + 0.3, y + 0.2, cw - 0.6, 0.35, "做什么", size=15, bold=True, color=ACC)
tb(s2, x + 0.3, y + 0.6, cw - 0.6, 0.8, "主持人与 AI 角色姬岁歌围绕游戏、弹幕和连续事件，共同制作直播节目", size=18, bold=True)
tb(s2, x + 0.3, y + 1.5, cw - 0.6, 0.8, [[("主形式  ", {"bold": True}), ("人机共播（计划书主线）", {})],
                                         [("AI 自主直播  ", {"bold": True}), ("定位为独立时段还是技术演示：首阶段验证方向待确认", {})]], size=14, color=MUTED, space=3)
# 2 首批观众
x, y = cards[1]; box(s2, x, y, cw, ch)
tb(s2, x + 0.3, y + 0.2, cw - 0.6, 0.35, "首批观众", size=15, bold=True, color=ACC)
tb(s2, x + 0.3, y + 0.6, cw - 0.6, 0.8, "已有 VUP 观看习惯、爱看游戏挑战与整活联动的中文观众", size=18, bold=True)
tb(s2, x + 0.3, y + 1.5, cw - 0.6, 0.8, [[("观看动机（待验证）  ", {"bold": True}), ("角色反差 · 临场互动 · 下一场回收约定", {})],
                                         [("已有接触  ", {"bold": True}), ("5 名路人初步口头反馈，非正式验证", {})]], size=14, color=MUTED, space=3)
# 3 当前进展
x, y = cards[2]; box(s2, x, y, cw, ch)
tb(s2, x + 0.3, y + 0.2, cw - 0.6, 0.35, "当前进展与证据", size=15, bold=True, color=ACC)
rows = [("已验收", "done", "Flappy Bird：100/100 局达 1,000 分"), ("局部验证", "wip", "STS2 18 次成功动作；Minecraft 木剑"),
        ("已有模块", "wip", "对话、记忆审核、语音、Live2D、工作台"), ("未开展", "plan", "整场直播、人机联机、正式观众验证")]
for i, (t, k, d) in enumerate(rows):
    yy = y + 0.65 + i * 0.42
    pill(s2, t, x + 0.3, yy, 1.05, k, size=11, h=0.3)
    tb(s2, x + 1.5, yy - 0.02, cw - 1.75, 0.34, d, size=14, anchor="m")
# 4 本次申请
x, y = cards[3]; box(s2, x, y, cw, ch, fill=ACCSOFT, line=None)
tb(s2, x + 0.3, y + 0.2, cw - 0.6, 0.35, "本次申请", size=15, bold=True, color=ACC)
tb(s2, x + 0.3, y + 0.55, cw - 0.6, 0.75, [[("46,200", {"size": 40, "bold": True, "color": ACC}), (" 元 · 6 个月", {"size": 18, "bold": True})]], anchor="b")
tb(s2, x + 0.3, y + 1.38, cw - 0.6, 0.95, [[("换来  ", {"bold": True}), ("原创形象交付 · 3 场内部试播 · 30 人观众验证 · B站主阵地试运营", {})],
                                           [("资金性质  ", {"bold": True}), ("支持性质待确认；不含团队工资", {})]], size=14, color=INK, space=3)
footer(s2, 2, "依据：商业计划书（2026-10-24）· demo-summary-20261018 · STS2 验收快照（2026-10-20）· 预算测算明细-改稿v2")
notes(s2, "【讲述重点】\n· 这一页回答评委最先关心的四件事：做什么、给谁看、现在做到哪、这笔钱换什么。\n· 主次关系：商业计划书以“主持人 + AI 共播”为主线；1025 稿还展示了 AI 自主直播。现有材料没有明确 AI 自主直播在首阶段的定位（独立时段或技术演示），页面标为“首阶段验证方向待确认”，需团队决定后改写。\n· 进展只列有记录的事实：Flappy Bird 独立验收（种子 30000–30099，100/100 局达到 1,000 分、零碰撞，来源 docs/demo-summary-20261018.md）；STS2 2026-10-20 快照 18 次成功动作、0 次工具错误（output/sts2/streamer-acceptance.json）；Minecraft 有采木、木剑等实机记录，但建家闭环未完成。Flappy Bird 尚未与主播语音完整集成。\n· 5 名路人反馈为口头初步反馈，没有书面记录前不扩大结论。\n· 申请额 46,200 元为团队测算（详见第 15、16 页），支持性质待确认，不含团队工资、税费和设备采购。\n【承接】下一页目录。", keep_old=False)

# ===================== 原第 2 页 → 新 03 目录 =====================
s = orig[1]
mods = [("项目背景", "首批观众为何观看并回访", "04"), ("市场规模与竞品分析", "已有 VUP 观众是切入口，差异待验证", "05"),
        ("产品介绍", "AI 自主直播 · MVP · 人机共播 · 观众验证 · 技术 · 关键问题", "06–11"),
        ("团队介绍", "核心成员与关键岗位", "12"), ("商业模式与获客", "先验证打赏与会员；B站主阵地与回访", "13–14"),
        ("市场战略与资金", "46,200 元预算与验收里程碑", "15–16")]
name_ids = [7, 12, 17, 22, 27, 32]; pg_ids = [9, 14, 19, 24, 29, 34]
for i, (a, b, pg) in enumerate(mods):
    nm = find(s, name_ids[i]); y = 0.75 + i * 0.98
    write(nm.text_frame, a, size=21, bold=True); move(nm, 4.97, y + 0.14, 5.9, 0.42)
    tb(s, 4.97, y + 0.57, 6.0, 0.32, b, size=14, color=MUTED)
    write(find(s, pg_ids[i]).text_frame, pg, size=16, bold=True, color=FAINT, align="r", anchor="m")
tb(s, 0.97, 3.0, 2.8, 0.7, [[("开篇", {"bold": True, "color": ACC})], [("项目摘要 · 02", {})]], size=15, color=MUTED, space=2)
pagenum(s, 3)
notes(s, "【增补修改版 · 讲述重点】\n· 目录按六个模块归组，页码与本文件一致：项目摘要 02；项目背景 04；市场与竞品 05；产品介绍 06–11（AI 自主直播、MVP 完成度、人机共播、观众验证、技术与角色资产、关键技术问题）；团队 12；商业模式与获客 13–14；市场战略、资金与里程碑 15–16。\n· 现场时间紧时，可直接跳过本页。")

# ===================== 原第 3 页 → 新 04 项目背景 =====================
s = orig[2]
write(find(s, 4).text_frame, "首批观众为何观看并回访姬岁歌", size=32, bold=True)
# 左侧三层
write(find(s, 6).text_frame, "目标观众", size=16, bold=True, color=ACC); move(find(s, 6), 0.85, 1.75)
write(find(s, 7).text_frame, ["已有 VUP 观看习惯", "爱看游戏挑战与搭档互动", "愿意追看角色的连续事件"], size=18); move(find(s, 7), 0.85, 2.2, 5.0, 1.1)
move(find(s, 5), 0.6, 1.75, 0, 1.55)
write(find(s, 9).text_frame, "观看动机（待验证）", size=16, bold=True, color=ACC); move(find(s, 9), 0.85, 3.6)
write(find(s, 10).text_frame, ["角色反差：岁好奇、歌坚持", "临场互动：主持人设局，AI 接梗", "下一场回收：约定与旧事"], size=18); move(find(s, 10), 0.85, 4.05, 5.0, 1.1)
move(find(s, 8), 0.6, 3.6, 0, 1.55)
line(s, 0.6, 5.45, 0.6, 6.75, color=LINE, w=3)
tb(s, 0.85, 5.45, 5.0, 0.4, "初步反馈", size=16, bold=True, color=ACC)
tb(s, 0.85, 5.9, 5.0, 0.85, "5 名路人观众偏好整活与联动内容（口头，初步定性，非正式验证）", size=16)
# 右侧节目循环：上移下排，补示例
write(find(s, 18).text_frame, "AI 角色回应", size=19, bold=True)
write(find(s, 19).text_frame, "按岁 / 歌性格作选择", size=15, color=MUTED)
for sid in (20, 21, 22, 23, 24, 25):
    sh = find(s, sid); sh.top = sh.top - I(0.75)
write(find(s, 21).text_frame, "游戏给出真实结果", size=19, bold=True)
write(find(s, 22).text_frame, "成功或翻车都在画面里", size=15, color=MUTED)
write(find(s, 24).text_frame, "观众参与并回访", size=19, bold=True)
write(find(s, 25).text_frame, "弹幕投票 · 下场回收", size=15, color=MUTED)
for sid in (26, 28):
    sh = find(s, sid)
    if sid == 28: sh.top = sh.top - I(0.75)
for sid in (27, 29):
    sh = find(s, sid); sh.height = I(0.9)
move(find(s, 30), y=3.3); move(find(s, 31), y=3.45)
ex = find(s, 32)
move(ex, 6.32, 5.6, 6.41, 1.25)
box(s, 6.32, 5.55, 6.41, 1.3, fill=ACCSOFT, line=None)
ex._element.getparent().remove(ex._element); s.shapes._spTree.append(ex._element)   # 文字置于底色之上
write(ex.text_frame, [[("节目设计示例  ", {"bold": True, "color": ACC}), ("主持人立规则“谁先做出木剑谁当队长”，AI 按角色决定先砍树还是先合成，弹幕投票下一步，下一场开播先清算队长约定。", {})],
                      [("待观测行为  ", {"bold": True, "color": ACC}), ("看完片段 · 主动互动 · 自愿回到下一场 · 记住角色与事件", {})]],
      size=13.5, color=INK, space=4)
move(ex, 6.5, 5.65, 6.08, 1.15)
pagenum(s, 4)
notes(s, "【增补修改版 · 讲述重点】\n· 标题改为观众视角：首批观众为何观看并回访。目标观众改为可识别的观看习惯：已有 VUP 观看习惯、爱看游戏挑战与搭档互动、愿意追看角色连续事件。年龄、平台偏好和付费经历需访谈补充，未自行编写。\n· 原稿“吸引观众、提升热度”改为待观测行为：是否看完片段、是否主动互动、是否自愿回到下一场、能否记住角色和事件。\n· 右下示例为节目设计，不是真实节目记录；“人化”含义见原备注（可识别的性格、表达与连续经历，不宣称真实情感）。原“人机分工”信息并入循环图的“主持人设局造梗 / AI 角色回应”。\n· 5 名路人反馈仅作初步定性接触，招募渠道、观看材料和具体意见待补（见素材清单）。\n【承接】下一页说明这群观众的已有习惯与竞争选择。")

# ===================== 原第 4 页 → 新 05 市场与竞品 =====================
s = orig[3]
write(find(s, 101).text_frame, "已有 VUP 观众是切入口，节目差异仍需真实观看验证", size=28.5, bold=True, color=INK)
tb(s, 0.6, 1.42, 12.13, 0.32, "中文、双角色、记忆、弹幕、游戏等能力同类项目已有，不作独有壁垒；比较维度统一为“观众看点”与“姬岁歌拟差异”", size=13.5, color=MUTED)
write(find(s, 107).text_frame, ["2025 年 CTR B站年轻人调研，N=4,056", "两项为不同问题的受访样本比例，不相乘、不外推人数"], size=11, color=FAINT)
move(find(s, 107), 0.6, 6.2, 3.8, 0.62)
cats = ["直接竞争 · 海外 AI 主播", "直接竞争 · 国内 AI 主播", "技术底座 · 开源工具", "替代内容 · 真人 VUP / AI 偶像"]
looks = ["成熟人物关系与长期内容", "双 AI 对话、记忆与游戏", "开源接入，降低技术门槛", "真人表演或 AI 角色故事"]
diffs = ["中文语境与北京晚间（待验证）", "真人主持 + 同体双人格（待验证）", "自营角色与连续节目，不卖工具", "即时选择与下场回收（待验证）"]
for i in range(4):
    write(find(s, 110 + i * 6).text_frame, cats[i], size=13.5, bold=True, color=ACC)
    fr = find(s, 108 + i * 6); fx, fy = fr.left / 914400, fr.top / 914400
    tb(s, fx + 0.21, fy + 1.2, 3.48, 0.3, [[("观众看点  ", {"bold": True, "color": MUTED}), (looks[i], {})]], size=13)
    lab = find(s, 112 + i * 6); write(lab.text_frame, "姬岁歌拟差异", size=12, bold=True, color=ACC); move(lab, y=fy + 1.6)
    d = find(s, 113 + i * 6); write(d.text_frame, diffs[i], size=14.5, color=INK); move(d, y=fy + 1.9, h=0.32)
write(find(s, 132).text_frame, [[("来源：[M1] CTR × B站《2026 年轻人消费趋势报告》", {"link": "https://www.ctrchina.cn/rich/report/802"}),
                                   ("    竞品：Vedal 官网、Lumi_Nox / AI-Vtuber / AI-YinMei 官方仓库、B站 Z 指数；口径见备注", {})]], size=10.5, color=FAINT)
pagenum(s, 5)
notes(s, "【增补修改版 · 讲述重点】\n· 左侧两项比例保留原样本说明：43.5% 为“在 B站浏览 VUP 直播或录播内容”、18.5% 为“曾为 VUP 打赏或充电”，均为 2025 年 CTR B站年轻人调研 N=4,056 的受访样本比例。两项是不同问题，不相乘、不相除，不乘平台月活外推人数，也不称为 AI 主播专属需求。\n· 四类对象分组：Neuro-sama、Lumi/Nox 为直接竞争（运营中的 AI 主播）；AI-Vtuber、吟美为开源技术底座（工具，不是运营角色）；传统 VUP 与 AI 偶像为替代内容（争夺同一批观众时间）。\n· 统一用“观众看点 / 姬岁歌拟差异”比较。原稿“基于庞大知识库的 AI 自主回应”缺乏区分度，改为“游戏中的即时选择与下场回收（待验证）”。\n· 页面上的竞品看点来自对方官网或创作者 README 的公开描述，不是对其运行效果的独立测试。\n【承接】下一页进入产品：先看 AI 自主直播的两个节目场景。")

# ===================== 原第 5 页 → 新 06 AI 自主直播 =====================
s = orig[4]
write(find(s, 3).text_frame, "03 产品介绍 · 1/6 AI 自主直播", size=13, bold=True, color=ACC)
write(find(s, 101).text_frame, "AI 自主直播：两个场景准备验证的节目看点", size=28.5, bold=True, color=INK)
write(find(s, 102).text_frame, "角色自主决策、游戏解说与弹幕互动，形成独立节目；首阶段定位（独立时段或技术演示）待确认", size=16, color=MUTED)
for px in (0.6, 6.88):
    pill(s, "节目概念示意 · 非实机", px + 0.15, 2.4, 2.35, "done", size=12, h=0.34)
write(find(s, 107).text_frame, [[("准备验证  ", {"bold": True, "color": ACC}), ("突发危险时，角色能否即时接梗并做出选择", {})]], size=15)
move(find(s, 107), 0.6, 6.2, 5.85, 0.6)
write(find(s, 108).text_frame, [[("准备验证  ", {"bold": True, "color": ACC}), ("出牌决策的反差能否引出吐槽与弹幕预测", {})]], size=15)
move(find(s, 108), 6.88, 6.2, 5.85, 0.6)
remove(find(s, 109))
write(find(s, 110).text_frame, "节目概念示意，非实机：AI 生成画面与拟写台词、弹幕，不代表已实现能力；两场景的观众观察项均为弹幕参与与回访", size=10.5, color=FAINT)
pagenum(s, 6)
notes(s, "【增补修改版 · 讲述重点】\n· 两张图保留，均为节目概念示意（AI 生成画面与拟写台词、弹幕），不是实机录屏或真实观众反馈；每张图左上角加了明显标识。\n· 每个场景改写为“准备验证的节目看点”：我的世界看突发危险时的即时接梗与选择；杀戮尖塔看出牌决策反差能否引出吐槽和弹幕预测。原底部“观众：弹幕参与与回访”并入页脚说明，作为两场景共同观察项。\n· AI 自主直播在首阶段是独立时段还是技术演示，材料未明确，待团队确认。\n【承接】下一页展示真正已经跑通的部分和尚待实测的指标。")

# ===================== 新 07 MVP 完成度 =====================
s7 = new_slide()
header(s7, "03 产品介绍 · 2/6 MVP 完成度", "MVP 已跑通局部链路，整场直播指标待实测")
iw = 6.55; ih = iw * 650 / 1510
s7.shapes.add_picture(ASSET_LIVE, I(0.6), I(1.55), I(iw), I(ih))
box(s7, 0.6, 1.55, iw, ih, fill=None, line=LINE, lw=0.75, rad=0)
pill(s7, "真实运行截图", 0.75, 1.68, 1.55, "done", size=12, h=0.32)
tb(s7, 0.6, 1.55 + ih + 0.06, iw, 0.5, "开发版游戏工作台，2026-10-14：AI 没做成工作台，口播“背包里只有云杉原木，得先把它变成木板”。Live2D 为临时占位模型，非姬岁歌最终形象", size=11.5, color=MUTED)
rx, rw = 7.45, 5.28
groups = [("已完成（有验收记录）", "done", ["Flappy Bird：100/100 局达 1,000 分，零碰撞", "STS2 快照：18 次成功动作，0 次工具错误"]),
          ("联调中", "wip", ["Minecraft 规划→执行：采木、木剑；建家未完成", "语音、字幕、Live2D 舞台与导演工作台"]),
          ("规划中", "plan", ["人类一败涂地联机控制", "B站 / 抖音真实事件接入；原创 Live2D 接入"])]
yy = 1.55
for t, k, items in groups:
    pill(s7, t, rx, yy, 2.2 if k == "done" else 1.2, k, size=12, h=0.32)
    for j, it in enumerate(items):
        tb(s7, rx + 0.15, yy + 0.42 + j * 0.36, rw - 0.15, 0.34, "· " + it, size=14)
    yy += 0.42 + len(items) * 0.36 + 0.18
fy = 5.08
tb(s7, 0.6, fy, 8, 0.35, "整场直播实测字段（同一版本、注明日期与条件）", size=15, bold=True, color=ACC)
fields = [("连续运行时长", "待实测", "开播到中断"), ("响应耗时", "待实测", "语音输入 → 首次可听回应"), ("任务成功 / 尝试", "18 / 18", "仅 STS2 快照；其余待实测"),
          ("人工接管", "待实测", "次数与原因"), ("调用成本", "待实测", "STS2 快照 57,826 tokens，费用未折算")]
fw = (12.13 - 4 * 0.2) / 5
for i, (a, v, d) in enumerate(fields):
    fx = 0.6 + i * (fw + 0.2)
    measured = v != "待实测"
    box(s7, fx, fy + 0.42, fw, 1.38, fill=WHITE, line=ACC if measured else FAINT, lw=1.25, dash=not measured)
    tb(s7, fx + 0.15, fy + 0.55, fw - 0.3, 0.3, a, size=13, bold=True, color=MUTED)
    tb(s7, fx + 0.15, fy + 0.85, fw - 0.3, 0.45, v, size=20, bold=True, color=ACC if measured else FAINT)
    tb(s7, fx + 0.15, fy + 1.32, fw - 0.3, 0.5, d, size=11, color=MUTED)
footer(s7, 7, "来源：docs/demo-summary-20261018.md · output/sts2/streamer-acceptance.json · logs/minecraft-deepseek-live-run.json")
notes(s7, "【讲述重点】\n· 这一页回答“现在到底做到了什么”。左侧是真实开发版截图（2026-10-14 游戏工作台）：AI 在 Minecraft 里没做成工作台，主动说明缺木板；Live2D 为临时占位模型，不是姬岁歌最终形象，也不是正式直播画面。\n· 已完成：Flappy Bird 独立策略验收，种子 30000–30099 共 100 局全部达到 1,000 分、零碰撞；另有单局 10,000 分加速耐久测试。该模块不调用大模型，尚未与主播语音完整集成。STS2 2026-10-20 快照记录 18 次成功动作、0 次工具错误（第 1 幕第 6 层），只是局部链路快照，不是整局或整场直播。\n· 联调中：Minecraft 有 DeepSeek / 千问规划 + Mineflayer 执行的实机记录，曾完成采木、木板、木棍、木剑等步骤；生存建家闭环未完成，复杂地形稳定性未充分证明。主播语音、字幕、Live2D 舞台、导演工作台为仓库已有模块，未做端到端重新验收。\n· 规划中：人类一败涂地联机控制、双平台真实事件接入（B站适配层目前基于模拟事件）、原创 Live2D 接入。\n· 实测字段：任务成功/尝试 18/18 仅指 STS2 快照（successfulActions=18，toolErrors=0）；调用量 57,826 tokens 来自同一快照，费用未折算。连续运行时长、响应耗时、人工接管均无记录，写“待实测”。\n【口径】响应耗时定义为“主持人语音输入结束 → 角色首次可听回应”，样本足够后再给 P50/P95。\n【承接】下一页回到节目：人机共播怎么设计、怎么实现。", keep_old=False)

# ===================== 原第 6 页 → 新 08 人机共播 =====================
s = orig[5]
write(find(s, 3).text_frame, "03 产品介绍 · 3/6 真人与 AI 联机共播", size=13, bold=True, color=ACC)
write(find(s, 101).text_frame, "人机共播的节目设计与实现路径", size=28.5, bold=True, color=INK)
write(find(s, 102).text_frame, [[("人类一败涂地 · 合作解谜节目构想  ", {}), ("联机能力待开发验证", {"bold": True, "color": ACC})]], size=16, color=MUTED)
pic = find(s, 103); ph = 4.25; pw = ph * 8.66 / 4.88
move(pic, 0.6, 2.1, pw, ph)
pill(s, "节目概念示意 · 非实机", 0.6 + pw - 2.5, 2.22, 2.35, "done", size=12, h=0.34)
tb(s, 0.6, 6.45, pw, 0.45, [[("实现路径  ", {"bold": True, "color": ACC}), ("先在已接入的 Minecraft 中验证主持人与 AI 的对话和任务协作，再开发本游戏联机；Minecraft 成果不能证明本游戏的控制能力", {})]], size=12.5)
rx = 0.6 + pw + 0.35; rw = 12.73 - rx
items = [("本场挑战", "限时把队友拉过断崖"), ("主持人任务", "定规则、出主意、给真实反应"), ("AI 角色任务", "按角色接话；联机动作待实现"),
         ("观众参与", "弹幕投票先救谁、怎么爬"), ("下场回收", "清算这场谁拖累了谁")]
pts = [(104, 105), (106, 107), (108, 109)]
for i, (a, b) in enumerate(items):
    y = 2.1 + i * 0.93
    if i < 3:
        pa, pb = find(s, pts[i][0]), find(s, pts[i][1])
        write(pa.text_frame, a, size=16, bold=True, color=ACC); move(pa, rx, y, rw, 0.34)
        write(pb.text_frame, b, size=15, color=INK); move(pb, rx, y + 0.36, rw, 0.5)
    else:
        tb(s, rx, y, rw, 0.34, a, size=16, bold=True, color=ACC)
        tb(s, rx, y + 0.36, rw, 0.5, b, size=15)
write(find(s, 110).text_frame, "节目效果示意，非实机录屏；对白与弹幕为拟写；人机联机控制尚待实现与验证", size=10.5, color=FAINT)
pagenum(s, 8)
notes(s, "【增补修改版 · 讲述重点】\n· 标题改为“节目设计与实现路径”；画面上明确标注节目概念示意、联机能力待开发验证。\n· 右侧改为一套可执行的节目结构（示例）：本场挑战、主持人任务、AI 角色任务、观众参与、下场回收。原“共同目标 / 双向互动 / 关系看点”的含义已并入这五项。\n· 实现路径：人类一败涂地的 AI 联机、抓取、攀爬尚未实现；更接近可演示的是已接入的 Minecraft，可先在其中验证主持人与 AI 的对话、主持和任务协作，再扩展新游戏。不同游戏的控制能力不能互相证明。\n· 人机共播是否比纯 AI 更值得额外人力，需在观众验证中比较（下一页）。\n【承接】下一页说明观众验证：已有什么、准备怎么测。")

# ===================== 新 09 观众验证 =====================
s9 = new_slide()
header(s9, "03 产品介绍 · 4/6 观众验证", "观众验证：已有 5 人初步反馈，30 人测试待执行")
lw = 5.1
box(s9, 0.6, 1.6, lw, 3.65)
pill(s9, "已有 · 初步定性", 0.9, 1.8, 1.75, "wip", size=12)
tb(s9, 0.9, 2.22, lw - 0.6, 0.45, "5 名路人观众", size=20, bold=True)
recs = [("主要反馈", "偏好整活与联动内容（团队口述）", True), ("招募渠道", "待补", False), ("观看材料", "待补", False), ("时间", "2026-10（具体日期待补）", False), ("据此修改", "待补", False)]
for i, (a, b, ok) in enumerate(recs):
    y = 2.8 + i * 0.47
    tb(s9, 0.9, y, 1.3, 0.4, a, size=14, bold=True, color=MUTED, anchor="m")
    tb(s9, 2.2, y, lw - 1.9, 0.4, b, size=14, color=INK if ok else FAINT, anchor="m")
rx = 0.6 + lw + 0.3; rw = 12.73 - rx
box(s9, rx, 1.6, rw, 3.65, line=FAINT, dash=True)
pill(s9, "计划 · 待执行", rx + 0.3, 1.8, 1.55, "plan", size=12)
tb(s9, rx + 0.3, 2.22, rw - 0.6, 0.45, "30 名成年受试者", size=20, bold=True)
tb(s9, rx + 0.3, 2.75, rw - 0.6, 0.3, "观察指标", size=14, bold=True, color=MUTED)
mets = ["观看时长", "角色辨识", "事件记忆", "主动回访", "付费意愿"]
mw = (rw - 0.6 - 4 * 0.12) / 5
for i, m in enumerate(mets):
    pill(s9, m, rx + 0.3 + i * (mw + 0.12), 3.1, mw, "soft", size=13, h=0.38)
tb(s9, rx + 0.3, 3.65, rw - 0.6, 1.5, [
    [("方法  ", {"bold": True, "color": ACC}), ("按观看习惯分层，随机平衡两段各 6 分钟样片的观看顺序", {})],
    [("探索性对比  ", {"bold": True, "color": ACC}), ("同主题、同时长的纯 AI 与人机共播短片", {})],
    [("预设门槛  ", {"bold": True, "color": ACC}), ("自由观看中位数 3 分钟；≥15 人记住角色或事件；≥6 人 7 天内主动回看满 60 秒", {})]], size=14, space=5)
ly = 5.5
tb(s9, 0.6, ly, 4, 0.35, "证据强度由弱到强（均未验证）", size=15, bold=True, color=ACC)
steps = [("口头意愿", "“愿意看 / 愿意买”"), ("真实回访", "自愿回到下一场"), ("真实付费", "平台净到账")]
sw = 3.55
for i, (a, b) in enumerate(steps):
    x = 0.6 + i * (sw + 0.75)
    box(s9, x, ly + 0.45, sw, 0.85, fill=SOFT if i == 0 else WHITE, line=LINE)
    tb(s9, x + 0.25, ly + 0.5, sw - 0.5, 0.4, a, size=17, bold=True)
    tb(s9, x + 0.25, ly + 0.88, sw - 0.5, 0.35, b, size=13, color=MUTED)
    if i < 2: line(s9, x + sw + 0.1, ly + 0.87, x + sw + 0.65, ly + 0.87, color=ACC, w=2, arrow=True)
footer(s9, 9, "依据：团队口述（5 人路人反馈，2026-10）；商业计划书第 19 页“观众验收与接梗复测”    小样本结果只用于改节目，不宣称已验证 PMF")
notes(s9, "【讲述重点】\n· 已有：5 名喜欢 VTuber、爱看整活与联动的路人观众的初步口头反馈，团队转述的主要意见是偏好整活与联动内容。没有书面记录，招募渠道、观看材料、具体日期和据此做的修改都待补。只能作为初步定性接触，不扩大结论。\n· 计划：30 名成年受试者，按观看习惯分层，随机平衡两段各 6 分钟样片观看顺序。观察观看时长、角色辨识、事件记忆、主动回访、付费意愿。预设门槛来自商业计划书：自由观看中位数 3 分钟、至少 15 人记住角色特点或事件、至少 6 人在 7 天内主动回看满 60 秒；报告保留有效分母、退出、缺失和技术故障。\n· 纯 AI 与人机共播的比较为探索性测试，小样本不做因果断言。\n· 证据强度：口头意愿 < 真实回访 < 真实付费。不把“愿意买”写成收入或付费转化，不宣称已验证 PMF。\n【承接】下一页回到技术方案与角色资产。", keep_old=False)

# ===================== 原第 7 页 → 新 10 技术与角色资产 =====================
s = orig[6]
write(find(s, 44).text_frame, "03 产品介绍 · 5/6 技术与角色资产", size=13, bold=True, color=ACC)
write(find(s, 2).text_frame, "直播角色引擎方案与周边概念", size=28.5, bold=True, color=INK)
write(find(s, 3).text_frame, "技术方案示意（标注实现状态）", size=18, bold=True, color=INK)
dashed = {14, 11, 26}   # 游戏状态、观众弹幕、游戏动作：联调或局部验证
for sid in (8, 11, 14, 17, 20, 23, 26):
    sh = find(s, sid)
    if sid in dashed:
        sh.line.dash_style = MSO_LINE.DASH
tb(s, 0.6, 6.62, 8.0, 0.3, [[("实线框  ", {"bold": True}), ("仓库已有模块（未做端到端重新验收）    ", {}), ("虚线框  ", {"bold": True}), ("联调中或局部验证", {})]], size=11.5, color=MUTED)
move(find(s, 38), y=6.3)
write(find(s, 40).text_frame, "立牌 · 徽章（概念设计，尚未打样）", size=14, bold=True, color=ACC)
write(find(s, 42).text_frame, "技术链路为方案示意，非已上线架构；周边为 AI 生成概念设计、尚未打样，无报价与销量", size=10.5, color=FAINT)
pagenum(s, 10)
notes(s, "【增补修改版 · 讲述重点】\n· 技术图与周边图均保留。技术节点按材料标注实现状态：实线框为仓库已有模块（ASR、人格与共同记忆、场景理解与调度、TTS/字幕/Live2D 表达），但本轮未做端到端重新验收；虚线框为联调中或局部验证（弹幕接入目前基于模拟事件；游戏状态与动作在 Minecraft / STS2 有局部实机记录）。整张图仍是方案示意，不是已上线架构。\n· 周边继续标注“概念设计，尚未打样”，无供应商报价、售价或销量。\n【承接】下一页把关键技术问题拆开：哪些是外部组件，哪些是团队工程，哪些是希望积累的资产。")

# ===================== 新 11 关键技术问题 =====================
s11 = new_slide()
header(s11, "03 产品介绍 · 6/6 关键技术问题", "关键技术问题与拟积累的资产")
cols = [("问题", 2.3), ("实现办法", 3.5), ("当前证据", 3.35), ("待解决", 2.98)]
y0 = 1.55; xs = [0.6]
for _, w in cols[:-1]: xs.append(xs[-1] + w)
for (t, w), x in zip(cols, xs):
    tb(s11, x + 0.12, y0, w - 0.2, 0.34, t, size=14, bold=True, color=ACC)
line(s11, 0.6, y0 + 0.4, 12.73, y0 + 0.4, color=ACC, w=1.25)
rows = [("说话时机", "主持人与 AI 不抢话", "语音识别送入上下文，调度决定谁、何时回应", "语音调度代码已有；首次可听回应耗时待实测", "实测延迟；长回复截断，主持人续场"),
        ("游戏执行", "说到做到", "模型规划 → 程序执行并检查结果 → 按准备 / 尝试 / 完成分别口播", "STS2 快照 18 次成功动作；Minecraft 采木与木剑", "长任务与复杂地形稳定性；建家闭环"),
        ("跨场记忆", "下一场能回收", "候选记忆经人工审核后写入 SQLite", "导演工作台记忆审核页已实现", "回收效果待观众验证；错误记忆纠正"),
        ("角色一致性", "岁与歌可区分", "人格设定与情绪状态模块约束回应", "角色状态模块已有；观众能否分辨待验证", "≥30 个中文场景测试，岁 / 歌分别复核")]
rh = 0.86
for i, (a, a2, b, c, d) in enumerate(rows):
    y = y0 + 0.5 + i * rh
    tb(s11, xs[0] + 0.12, y, cols[0][1] - 0.2, rh - 0.1, [[(a, {"bold": True, "size": 16})], [(a2, {"size": 12.5, "color": MUTED})]], size=16)
    tb(s11, xs[1] + 0.12, y, cols[1][1] - 0.25, rh - 0.1, b, size=13.5)
    tb(s11, xs[2] + 0.12, y, cols[2][1] - 0.25, rh - 0.1, c, size=13.5, color=ACC if "已" in c[:12] else INK)
    tb(s11, xs[3] + 0.12, y, cols[3][1] - 0.2, rh - 0.1, d, size=13.5, color=MUTED)
    line(s11, 0.6, y + rh - 0.06, 12.73, y + rh - 0.06, color=LINE, w=0.75)
by = 5.45; bw = (12.13 - 2 * 0.25) / 3
blocks = [("外部模型与组件", "可替换", "SenseVoice · DeepSeek / 千问 · GPT-SoVITS · Live2D · OBS · Mineflayer", SOFT, None, False),
          ("团队工程工作", "已有 / 联调中", "调度与人工接管 · 游戏桥接与结果检查 · 记忆审核 · 导演工作台", ACCSOFT, None, False),
          ("拟积累的资产", "未验证", "角色设定与声线 · 节目事件库 · 中文接梗场景库 · 切片素材", WHITE, FAINT, True)]
for i, (a, t, b, f, l, dsh) in enumerate(blocks):
    x = 0.6 + i * (bw + 0.25)
    box(s11, x, by, bw, 1.38, fill=f, line=l, dash=dsh)
    tb(s11, x + 0.25, by + 0.15, bw - 0.5, 0.38, [[(a + "  ", {"bold": True, "size": 16}), (t, {"size": 12, "color": MUTED})]], size=16)
    tb(s11, x + 0.25, by + 0.6, bw - 0.5, 0.7, b, size=13.5)
footer(s11, 11, "依据：商业计划书第 8–9 页；工程索引 [I1][I2][I4][I6]；docs/demo-summary-20261018.md    未验证的只称拟建立优势，不宣称技术领先或壁垒")
notes(s11, "【讲述重点】\n· 用“问题 — 实现办法 — 当前证据 — 待解决”拆开四个关键问题：说话时机、游戏执行、跨场记忆、角色一致性。\n· 说话时机：SenseVoice 把主持人语音送入上下文，调度决定谁、何时回应；仓库有语音调度代码，但首次可听回应耗时没有实测记录。\n· 游戏执行：模型规划后由程序执行并读取结果，口播按准备、尝试、完成区分，避免把计划说成成果。证据为 STS2 2026-10-20 快照与 Minecraft 实机记录；长任务稳定性与建家闭环未完成。\n· 跨场记忆：候选记忆经人工确认后写入 SQLite（src/brain/memory-candidates.js、studio.html 记忆审核页）；能否让观众感到“被记住”需观众验证。\n· 角色一致性：岁偏好奇观察、歌偏坚持守护；观众能否分辨两种性格待测试，计划用不少于 30 个中文场景，每种人格分别生成并复核。\n· 底部三栏区分：外部模型与组件可替换；团队自己的工程工作是调度、接管、游戏桥接、记忆审核和工作台；希望积累的角色、事件、接梗场景与切片素材属于拟建立资产，尚未验证。\n【承接】下一页介绍完成这些工作的团队。", keep_old=False)

# ===================== 原第 8 页 → 新 12 团队 =====================
s = orig[7]
write(find(s, 4).text_frame, "核心团队与关键岗位", size=32, bold=True)
tb(s, 0.6, 1.36, 12.13, 0.38, "当前 46,200 元预算不含团队工资：成员暂不领薪，工时单独记录；履历、作品与投入时间由团队补充", size=15, color=MUTED)
for sid in (14, 15, 16, 17, 24, 37):
    remove(find(s, sid))
core = [(5, 6, 7, "项目统筹与资源决策"), (8, 9, 10, "模型、游戏与系统验收"), (11, 12, 13, "节目排期与团队协作")]
cwid = 7.35
for i, (b, n, d, dl) in enumerate(core):
    y = 1.9 + i * 1.62
    bx = find(s, b); move(bx, 0.6, y, cwid, 1.48)
    move(find(s, n), 0.85, y + 0.12, cwid - 0.5, 0.42)
    dd = find(s, d); move(dd, 0.85, y + 0.52, cwid - 0.5, 0.34)
    col = WHITE if i == 0 else MUTED
    write(dd.text_frame, [[("本项目交付  ", {"bold": True}), (dl, {})]], size=14, color=WHITE if i == 0 else INK)
    fw3 = (cwid - 0.5 - 2 * 0.12) / 3
    for j, lab in enumerate(["相关经历", "实际作品", "身份 / 每周投入"]):
        fx = 0.85 + j * (fw3 + 0.12)
        box(s, fx, y + 0.92, fw3, 0.44, fill=None, line=ACC3 if i == 0 else FAINT, lw=1, dash=True)
        tb(s, fx + 0.12, y + 0.92, fw3 - 0.2, 0.44, [[(lab + "  ", {"bold": True}), ("待补", {})]], size=12.5, color=WHITE if i == 0 else FAINT, anchor="m")
rx = 0.6 + cwid + 0.35; rw = 12.73 - rx
box(s, rx, 1.9, rw, 1.42, fill=ACCSOFT, line=None)
tb(s, rx + 0.22, 1.98, rw - 0.4, 0.32, "关键岗位", size=14, bold=True, color=ACC)
tb(s, rx + 0.22, 2.32, rw - 0.4, 0.95, [[("主持  ", {"bold": True}), ("按场次轮值，人选与连续搭档安排待补", {})],
                                        [("技术值守与人工接管  ", {"bold": True}), ("徐嘉盛（计划书）", {})],
                                        [("节目决策  ", {"bold": True}), ("姬祥排期 · 杨若朝投入决策", {})]], size=12.5, space=2)
others = [(18, 19, 20), (21, 22, 23), (25, 26, 27), (28, 29, 30), (31, 32, 33), (34, 35, 36)]
for i, (b, n, d) in enumerate(others):
    y = 3.45 + i * 0.58
    move(find(s, b), rx, y, rw, 0.52)
    nm = find(s, n); txt = nm.text_frame.text
    name, role = txt.split("  ", 1) if "  " in txt else (txt, "")
    write(nm.text_frame, [[(name + "  ", {"bold": True, "size": 14}), (role, {"size": 11.5, "bold": True, "color": ACC})]], size=14)
    move(nm, rx + 0.18, y + 0.03, rw - 0.3, 0.26)
    dd = find(s, d); write(dd.text_frame, dd.text_frame.text, size=11.5, color=MUTED); move(dd, rx + 0.18, y + 0.27, rw - 0.3, 0.22)
pagenum(s, 12)
notes(s, "【增补修改版 · 讲述重点】\n· 九名成员与分工全部保留；组织层级连线去掉，降低头衔比重。左侧突出三位核心成员（按计划书岗位暂定为杨若朝、徐嘉盛、姬祥，若实际承担关键工作的成员不同请替换），每人留出相关经历、实际作品、身份与每周投入三个字段，材料未提供，写“待补”，不代写履历。\n· 关键岗位：主持按场次轮值，人选和保持搭档关系连续性的安排待补；技术值守与人工接管按计划书由技术总监负责依赖服务和接管流程；节目方向由运营总监排期、CEO 决定下一轮投入。\n· 学生 / 兼职 / 全职状态未提供，待补。\n· 当前 46,200 元预算不含团队工资，成员暂不领薪、工时单独记录；长期人力安排需另行说明。\n【承接】下一页讲收入顺序和单场经营模型。")

# ===================== 原第 9 页 → 新 13 商业模式 =====================
s = orig[8]
write(find(s, 3).text_frame, "05 商业模式与获客 · 1/2", size=13, bold=True, color=ACC)
write(find(s, 101).text_frame, "先验证打赏与会员，再拓展周边和商务", size=28.5, bold=True, color=INK)
write(find(s, 103).text_frame, "近期验证 · 首阶段", size=17, bold=True, color=ACC); move(find(s, 103), y=1.72)
write(find(s, 104).text_frame, "条件满足后拓展", size=17, bold=True, color=MUTED); move(find(s, 104), y=1.72)
move(find(s, 105), y=2.16); move(find(s, 106), y=2.16)
objs = ["直播观众", "核心粉丝", "喜欢角色的观众", "品牌与游戏厂商"]
conds = ["主阵地公开试运营起", "出现稳定回访观众后开通测试", "有明确需求、可交付后小批量预售", "有受众画像与节目样例后测试"]
for i in range(4):
    base = 107 + i * 5 - (1 if i else 0)
    num = find(s, [107, 112, 117, 122][i]); inc = find(s, [108, 113, 118, 123][i])
    met = find(s, [110, 115, 120, 125][i]); val = find(s, [111, 116, 121, 126][i])
    x = num.left / 914400
    move(num, y=2.3); move(inc, y=2.68); move(met, y=3.2); move(val, y=3.62, h=0.7)
    tb(s, x, 4.42, 2.69, 0.3, [[("付费对象  ", {"bold": True, "color": MUTED}), (objs[i], {})]], size=13.5)
    tb(s, x, 4.78, 2.69, 0.6, [[("启动条件  ", {"bold": True, "color": MUTED}), (conds[i], {})]], size=13.5)
move(find(s, 127), y=5.5)
uy = 5.62
tb(s, 0.6, uy, 6, 0.32, "单场经营模型（每场节目，待实测）", size=15, bold=True, color=ACC)
ue = [("净到账收入", "打赏 + 会员，扣平台分成"), ("模型语音成本", "LLM / TTS / ASR 调用"), ("主持与剪辑工时", "按小时登记，不计现金"), ("其他直接成本", "推广、素材、杂费"), ("单场贡献", "收入 − 直接现金成本")]
uw = (12.13 - 4 * 0.15) / 5
for i, (a, d) in enumerate(ue):
    x = 0.6 + i * (uw + 0.15)
    box(s, x, uy + 0.4, uw, 0.88, fill=ACCSOFT if i == 4 else WHITE, line=None if i == 4 else FAINT, dash=i != 4)
    tb(s, x + 0.15, uy + 0.46, uw - 0.3, 0.3, [[(a + "  ", {"bold": True}), ("待实测" if i < 4 else "", {"color": FAINT})]], size=13.5)
    tb(s, x + 0.15, uy + 0.8, uw - 0.3, 0.42, d, size=11.5, color=MUTED)
write(find(s, 132).text_frame, "四类收入为规划，尚未形成经营实绩或商业闭环；历史财务测算（预算v2，2026-10-24）仅留备注备答", size=10.5, color=FAINT)
pagenum(s, 13)
notes(s, "【增补修改版 · 讲述重点】\n· 四类收入全部保留，按开展顺序分为“近期验证”（打赏、会员）与“条件满足后拓展”（周边、商务），并补付费对象与启动条件。\n· 底部单场经营模型：净到账收入（打赏 + 会员，扣除平台分成）、模型语音成本、主持与剪辑工时（登记工时，不计入现金）、其他直接成本，以及单场贡献 = 收入 − 直接现金成本。目前没有开播，全部写“待实测”。\n· 原备注中的财务数字属于历史版本情景（预算测算明细-改稿v2.xlsx，2026-10-24）：全年收入情景 61,650 元、运营与履约支出 62,500 元，差额 −850 元，未计启动投入、团队工资与税费。该口径下收入只覆盖部分成本，不能讲成已盈利或已形成商业闭环；本版主页面不展示这些数字，评委问到时按历史情景回答。\n【承接】下一页说明观众从哪里来、怎么回访。")

# ===================== 新 14 获客与回访 =====================
s14 = new_slide()
header(s14, "05 商业模式与获客 · 2/2", "获客与回访：B站为主阵地，抖音做轻量测试")
lw = 6.3
ch_rows = [("B站", "主阵地", "直播 + 连载，承接追更、弹幕与社群", "done"), ("抖音", "试验渠道", "切片分发 + 低频直播测试，按效果决定投入", "wip"),
           ("切片", "计划每周 4–5 条", "保留抛梗、回应与结局，导流到直播", "soft"), ("社群", "粉丝群 / 动态", "预告下一场，提醒回收约定", "soft")]
for i, (a, t, d, k) in enumerate(ch_rows):
    y = 1.6 + i * 1.05
    box(s14, 0.6, y, lw, 0.92)
    tb(s14, 0.85, y + 0.12, 1.2, 0.4, a, size=19, bold=True)
    pill(s14, t, 0.85, y + 0.55, 1.75 if len(t) > 5 else 1.2, k, size=11, h=0.28)
    tb(s14, 2.95, y, lw - 2.55, 0.92, d, size=15, anchor="m")
rx = 0.6 + lw + 0.45; rw = 12.73 - rx
tb(s14, rx, 1.6, rw, 0.35, "切片到回访的漏斗（计划指标，均待实测）", size=15, bold=True, color=ACC)
fun = ["切片曝光", "有效观看", "关注", "直播到访", "下一场回访"]
for i, f in enumerate(fun):
    w = rw - i * 0.55; x = rx + (rw - w) / 2; y = 2.05 + i * 0.66
    box(s14, x, y, w, 0.56, fill=ACC if i == 4 else (ACCSOFT if i % 2 == 0 else RGBColor(0xDC, 0xE5, 0xFA)), line=None, rad=0.1)
    tb(s14, x, y, w, 0.56, [[(f + "  ", {"bold": True, "color": WHITE if i == 4 else INK}), ("待实测", {"size": 12, "color": WHITE if i == 4 else FAINT})]], size=16, align="c", anchor="m")
box(s14, 0.6, 5.95, 12.13, 0.8, fill=SOFT, line=None)
tb(s14, 0.85, 5.95, 11.7, 0.8, [[("已有结果  ", {"bold": True, "color": ACC}), ("尚无公开运营数据。", {}), ("以上渠道分工与每周 4–5 条切片为计划；指标定义：同期同口径统计单条曝光、有效播放、关注与可追踪的直播到访，不用总播放替代回访。", {"color": MUTED})]], size=14, anchor="m")
footer(s14, 14, "依据：商业计划书第 17–18 页（双平台直播方案）；1025 稿第 10 页渠道图；评审方案（2026-10-24）")
notes(s14, "【讲述重点】\n· 渠道分工：B站为主阵地（直播与连载，承接追更、弹幕与社群）；抖音为试验渠道（切片分发与低频直播测试，按效果决定是否加大投入）；切片计划每周 4–5 条；社群负责预告下一场和回收约定。\n· 调整说明：计划书写的是 B站、抖音双平台直播；考虑到团队产能尚未验证，本版改为“主阵地 + 试验渠道”，不默认能同时高强度经营两端。该渠道主次已由团队确认（2026-10-25）；是否加大抖音直播，以试运营数据决定。\n· 漏斗：切片曝光 → 有效观看 → 关注 → 直播到访 → 下一场回访，全部为计划指标，没有已有结果。统计口径：同一时期、同一口径记录单条内容数据和可追踪的直播到访，不用总播放替代回访。\n【承接】下一页是资金预算。", keep_old=False)

# ===================== 原第 10 页 → 新 15 市场战略与资金 =====================
s = orig[9]
write(find(s, 3).text_frame, "06 市场战略与资金 · 1/2", size=13, bold=True, color=ACC)
write(find(s, 4).text_frame, "46,200 元支持六个月原型与试播验证", size=32, bold=True)
write(find(s, 9).text_frame, "主阵地 · 连载", size=12, color=MUTED, align="c")
write(find(s, 14).text_frame, "试验 · 切片", size=12, color=MUTED, align="c")
write(find(s, 25).text_frame, "形象交付 → 内部试播与接梗测试 → 观众验证 → B站开播、抖音轻量测试", size=13, color=INK)
write(find(s, 22).text_frame, "首阶段 · 到账后第 1–6 月", size=15.5, bold=True, color=ACC)
write(find(s, 27).text_frame, "第 7–12 月", size=15.5, bold=True, color=INK)
write(find(s, 29).text_frame, "依靠到账收入", size=12, bold=True, color=MUTED)
write(find(s, 30).text_frame, "固定栏目 → 会员测试 → 小批量周边 → 品牌样例；收入未到账则延后", size=13, color=INK)
lab = find(s, 32); move(lab, 0.6, 4.35, 5.2, 0.35); write(lab.text_frame, "首阶段申请支持（支持性质待确认）", size=15, bold=True, color=MUTED)
big = find(s, 33); move(big, 0.6, 4.7, 5.2, 0.8)
write(big.text_frame, [[("46,200", {"size": 48, "bold": True, "color": ACC}), (" 元", {"size": 20})]], anchor="b")
grp = [("启动 25,000", "原画 8,000 · Live2D 14,000 · 工具 3,000"), ("运营 17,000", "语音云 9,000 · 剪辑 3,000 · 推广 3,000 · 杂费 2,000"), ("预备金 4,200", "支出小计的 10%")]

for i, (a, b) in enumerate(grp):
    y = 5.6 + i * 0.4
    tb(s, 0.6, y, 5.35, 0.38, [[(a + "  ", {"bold": True, "color": INK}), (b, {"color": MUTED})]], size=12, anchor="m")
write(find(s, 35).text_frame, [[("资金用途（元）", {"bold": True}), ("  不含团队工资 · 税费未单列 · 设备采购未列（待确认）", {"size": 11.5, "color": MUTED})]], size=16)
write(find(s, 38).text_frame, [[("预算来源：团队测算及报价依据（原画、Live2D 报价单待附）；明细见《预算测算明细-改稿v2.xlsx》第 5 表    渠道接入参考 ", {}),
                                 ("[S8] 抖音开放平台直播 SDK", {"color": ACC, "link": "https://open.douyin.com/platform/resource/docs/ability/douyin-live-sdk/introduction"})]], size=10.5, color=FAINT)
pagenum(s, 15)
notes(s, "【增补修改版 · 讲述重点】\n· 原渠道图、阶段计划、申请金额和八项预算全部保留，重新分层：启动 25,000（原画 8,000、Live2D 及表情 14,000、开发工具 3,000）；精简运营 17,000（模型语音云服务 9,000、剪辑补贴 3,000、推广测试 3,000、活动杂费 2,000）；预备金 4,200（支出小计的 10%）。合计 46,200 元，八项加总已复核。\n· 原画与 Live2D 合计 22,000 元，占申请额约 47.6%：观众要记住的是角色，节目需要可持续使用的原创形象，目前的 Live2D 为临时占位模型，不能用于正式播出；两项尚未付款，计划 2026 年 11 月底交付，建议按交付物验收后付款（付款节点待与画师确认）。若观众验证不佳，后续表情、动作和周边设计不追加。\n· 预算边界：不含团队工资（成员暂不领薪）；税费未单列；设备采购未列入，是否使用成员自有设备待确认。支持性质（赛事支持、补助、赞助或其他）待确认，不设置估值或股权条款。\n· 时间统一为“资金到账后第 X 月”。渠道图中 B站标为主阵地、抖音标为试验渠道，与上一页一致。\n· 来源改为“团队测算及报价依据”；抖音开放平台直播 SDK 只作为渠道接入参考，不支持预算数字。\n【承接】最后一页说明每个阶段交付什么、什么情况下继续或调整。")

# ===================== 新 16 验收里程碑 =====================
s16 = new_slide()
header(s16, "06 市场战略与资金 · 2/2", "验收里程碑：用运行与回访证据决定是否继续投入")
ms = [("到账后第 1 月", "形象交付", "原画与 Live2D 文件、表情与口型测试录屏"), ("第 2 月", "内部试播", "3 场试播录像；运行日志（时长、接管、成本）；≥10 个切片候选"),
      ("第 3 月", "观众验证", "30 人测试记录：观看时长、角色与事件记忆、7 天回访"), ("第 4–6 月", "主阵地试运营", "B站开播记录、回访与净到账收入台账；抖音轻量测试数据")]
mw = (12.13 - 3 * 0.25) / 4; my = 1.6
line(s16, 0.6 + 0.21, my + 0.21, 0.6 + 3 * (mw + 0.25) + 0.21, my + 0.21, color=ACC3, w=2)
for i, (t, a, b) in enumerate(ms):
    x = 0.6 + i * (mw + 0.25)
    oval = s16.shapes.add_shape(MSO_SHAPE.OVAL, I(x), I(my), I(0.42), I(0.42))
    oval.fill.solid(); oval.fill.fore_color.rgb = ACC; oval.line.fill.background(); oval.shadow.inherit = False
    tb(s16, x, my, 0.42, 0.42, str(i + 1), size=14, bold=True, color=WHITE, align="c", anchor="m")
    tb(s16, x + 0.55, my + 0.04, mw - 0.55, 0.34, t, size=14, bold=True, color=ACC)
    box(s16, x, my + 0.6, mw, 1.75)
    tb(s16, x + 0.22, my + 0.75, mw - 0.4, 0.4, a, size=19, bold=True)
    tb(s16, x + 0.22, my + 1.2, mw - 0.4, 1.1, [[("验收物  ", {"bold": True, "color": MUTED}), (b, {})]], size=13.5)
tb(s16, 0.6, my + 2.45, 12.13, 0.3, "原计划 2026 年 11 月底交付形象，进度取决于画师排期；以上统一按资金到账后月份计", size=12, color=FAINT)
cy = 4.55; hw = (12.13 - 0.3) / 2
box(s16, 0.6, cy, hw, 1.55, fill=ACCSOFT, line=None)
tb(s16, 0.85, cy + 0.15, hw - 0.5, 0.35, "继续投入的证据", size=16, bold=True, color=ACC)
tb(s16, 0.85, cy + 0.55, hw - 0.5, 0.95, ["节目能稳定播完，人工接管可控", "观众记住角色与事件，并有自愿回访", "净到账收入能覆盖单场直接现金成本（不含工资）"], size=14, space=3)
box(s16, 0.6 + hw + 0.3, cy, hw, 1.55, line=FAINT, dash=True)
tb(s16, 0.85 + hw + 0.3, cy + 0.15, hw - 0.5, 0.35, "需要调整时", size=16, bold=True)
tb(s16, 0.85 + hw + 0.3, cy + 0.55, hw - 0.5, 0.95, [[("回访不足  ", {"bold": True}), ("先改节目，暂停追加推广", {})],
                                                     [("收入未到账  ", {"bold": True}), ("延后周边、新游戏接入与抖音直播加量", {})],
                                                     [("始终保留  ", {"bold": True}), ("模型语音服务、主阵地开播、核心测试", {})]], size=14, space=3)
box(s16, 0.6, 6.25, 12.13, 0.6, fill=ACC, line=None)
tb(s16, 0.85, 6.25, 11.7, 0.6, [[("本次申请 46,200 元  ", {"bold": True, "size": 17}), ("完成形象交付、内部试播、30 人观众验证与 B站主阵地试运营（支持性质待确认）", {})]], size=14.5, color=WHITE, anchor="m")
footer(s16, 16, "依据：预算测算明细-改稿v2.xlsx 第 5 表；商业计划书第 18–19 页（阶段目标与验收）；评审方案（2026-10-24）")
notes(s16, "【讲述重点】\n· 四个阶段各对应验收物，时间统一按资金到账后月份：第 1 月形象交付（原画与 Live2D 文件、表情与口型测试录屏）；第 2 月内部试播（3 场录像、运行日志，以及至少 10 个切片候选，来自计划书首月目标）；第 3 月观众验证（30 人测试记录）；第 4–6 月 B站主阵地试运营（开播记录、回访与净到账收入台账，抖音轻量测试）。\n· 完成事项不等于验证成功。继续投入看三条证据：节目能稳定播完且人工接管可控；观众记住角色与事件并自愿回访；净到账收入覆盖单场直接现金成本（不含工资）。\n· 需要调整时：回访不足先改节目、暂停追加推广；收入未按计划到账，延后周边、新游戏接入（如人类一败涂地联机）和抖音直播加量；始终保留模型语音服务、主阵地开播和核心测试。\n· 原画与 Live2D 原计划 2026 年 11 月底交付，进度依赖画师排期。\n· 收尾：本次申请 46,200 元，用于完成上述阶段验证；支持性质待确认，不涉及估值、股份或其他融资承诺。", keep_old=False)

# ---------- 重新排序 ----------
order = [orig[0], s2, orig[1], orig[2], orig[3], orig[4], s7, orig[5], s9, orig[6], s11, orig[7], orig[8], s14, orig[9], s16]
sldIdLst = prs.slides._sldIdLst
id_el = {int(el.get("id")): el for el in sldIdLst}
order_ids = [sl.slide_id for sl in order]
for el in list(sldIdLst):
    sldIdLst.remove(el)
for sid in order_ids:
    sldIdLst.append(id_el[sid])
prs.save(OUT)
print("saved", OUT, len(prs.slides))
