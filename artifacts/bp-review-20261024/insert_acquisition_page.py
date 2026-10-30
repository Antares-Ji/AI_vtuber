"""把《增补修改版-v2》第 14 页“获客与回访”原样插入原版《1025稿》，放在原第 9 页（商业模式）之后、原第 10 页（市场战略）之前。
原版其余 10 页不做任何修改。
用法：py -3 insert_acquisition_page.py <1025稿.pptx> <增补修改版-v2.pptx> <输出.pptx>"""
import sys, copy
from pptx import Presentation

ORIG, SRC, OUT = sys.argv[1:4]
dst = Presentation(ORIG)
src_slide = Presentation(SRC).slides[13]

new = dst.slides.add_slide(dst.slide_layouts[0])
for ph in list(new.placeholders):
    ph._element.getparent().remove(ph._element)
new.background.fill.solid()
new.background.fill.fore_color.rgb = src_slide.background.fill.fore_color.rgb
tree = new.shapes._spTree
for sh in src_slide.shapes:                         # 本页只有文本框、形状和线条，无图片与外部链接
    tree.append(copy.deepcopy(sh._element))

for sh in new.shapes:
    if not sh.has_text_frame:
        continue
    t = sh.text_frame.text
    if t == "05 商业模式与获客 · 2/2":              # 原版第 9 页栏目名为“05 商业模式”，去掉“2/2”以免对不上
        sh.text_frame.paragraphs[0].runs[0].text = "05 商业模式 · 获客与回访"
        for r in sh.text_frame.paragraphs[0].runs[1:]:
            r.text = ""
    elif t.strip().endswith("/ 16"):                # 原版页码本身不连续（/10 与 /11 混用），此页不放页码
        sh._element.getparent().remove(sh._element)

# 讲者备注
notes_src = src_slide.notes_slide.notes_text_frame.text
ns = new.notes_slide
if ns.notes_text_frame is None:
    ns.shapes._spTree.append(copy.deepcopy(dst.slides[0].notes_slide.notes_placeholder._element))
ns.notes_text_frame.text = notes_src

# 移到原第 9 页之后（第 10 张）
lst = dst.slides._sldIdLst
el = lst[-1]
lst.remove(el)
lst.insert(9, el)
dst.save(OUT)
print("saved", OUT, len(dst.slides))
